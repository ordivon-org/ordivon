#!/usr/bin/env python3
"""Mechanical smoke doctor for the replaceable Artifact Build & Delivery toolchain.

The doctor checks local tool availability, exact selected versions, shared profile
contracts, deterministic OOXML generation/validation, and primary local web
runners. It deliberately does not claim Microsoft Office target acceptance,
visual quality, PDF/UA conformance of arbitrary files, or delivery completion.
"""
from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import artifact_trust.vsa as artifact_trust_vsa
from artifact_core.profile_v1 import validate_profile_v1
from artifact_operations.providers import DirectPythonOperationProvider
from scripts.artifact_oci_package import execute_oci_package_stage

GLOBAL_ARTIFACT_TOOLCHAIN_ROOT = Path(os.environ.get("ARTIFACT_TOOLCHAIN_ROOT", "/opt/ordivon/external/artifact-toolchain"))
GLOBAL_PANDOC = GLOBAL_ARTIFACT_TOOLCHAIN_ROOT / "pandoc/3.10.2/bin/pandoc"
GLOBAL_VERAPDF = GLOBAL_ARTIFACT_TOOLCHAIN_ROOT / "verapdf/1.30.2/verapdf"
GLOBAL_VNU = GLOBAL_ARTIFACT_TOOLCHAIN_ROOT / "vnu/26.9.7/vnu.jar"
GLOBAL_COSIGN = GLOBAL_ARTIFACT_TOOLCHAIN_ROOT / "cosign/3.1.3/bin/cosign"
GLOBAL_NODE_PACKAGE_ROOT = GLOBAL_ARTIFACT_TOOLCHAIN_ROOT / "node/1.63.0"
LOCK = json.loads((ROOT / "artifact-delivery/toolchain-v1.lock.json").read_text())
OPENXML_LOCK = json.loads((ROOT / "artifact-delivery/openxml-runtime-v1.lock.json").read_text())


def run(
    command: list[str],
    *,
    cwd: Path = ROOT,
    timeout: int = 60,
    env: dict[str, str] | None = None,
) -> subprocess.CompletedProcess[str]:
    """Run one diagnostic probe without letting an unavailable tool crash Doctor."""
    try:
        return subprocess.run(
            command,
            cwd=cwd,
            text=True,
            capture_output=True,
            timeout=timeout,
            check=False,
            env=env,
        )
    except FileNotFoundError as error:
        return subprocess.CompletedProcess(
            command,
            127,
            stdout="",
            stderr=f"executable unavailable: {error}",
        )
    except subprocess.TimeoutExpired as error:
        stdout = error.stdout if isinstance(error.stdout, str) else ""
        stderr = error.stderr if isinstance(error.stderr, str) else ""
        return subprocess.CompletedProcess(
            command,
            124,
            stdout=stdout,
            stderr=(stderr + f"\nprobe timed out after {timeout}s").lstrip(),
        )


def record(checks: list[dict[str, Any]], name: str, ok: bool, **details: Any) -> None:
    checks.append({"name": name, "status": "PASS" if ok else "FAIL", **details})


def selected_path(env_name: str, global_default: Path, legacy_default: Path | None = None) -> str:
    configured = os.environ.get(env_name)
    if configured:
        return configured
    if global_default.exists():
        return str(global_default)
    if legacy_default is not None and legacy_default.exists():
        return str(legacy_default)
    return str(global_default)


def main() -> int:
    checks: list[dict[str, Any]] = []
    provider = DirectPythonOperationProvider()

    python = selected_path("ARTIFACT_PYTHON", Path("/root/.local/share/ordivon-workstation/artifact-delivery-python-v1/current/bin/python"))
    snippet = (
        "import json; from importlib.metadata import version; "
        "print(json.dumps({p:version(p) for p in "
        "['jsonschema','python-pptx','XlsxWriter','opentelemetry-sdk']}))"
    )
    proc = run([python, "-c", snippet])
    if proc.returncode == 0:
        observed = json.loads(proc.stdout)
        expected = {
            "jsonschema": LOCK["python"]["jsonschema"],
            "python-pptx": LOCK["python"]["python-pptx"],
            "XlsxWriter": LOCK["python"]["XlsxWriter"],
            "opentelemetry-sdk": LOCK["python"]["opentelemetry-sdk"],
        }
        record(checks, "python-packages", observed == expected, observed=observed, expected=expected)
    else:
        record(checks, "python-packages", False, stderr=proc.stderr[-2000:])

    profile_results: dict[str, Any] = {}
    profile_ok = True
    profile_candidates = []
    for candidate in sorted((ROOT / "artifact-delivery/examples").glob("*.json")):
        try:
            value = json.loads(candidate.read_text())
        except Exception:
            continue
        if isinstance(value, dict) and value.get("profileVersion") == 1 and isinstance(value.get("artifactClass"), str):
            profile_candidates.append(candidate)
    for profile in profile_candidates:
        value = validate_profile_v1(profile)
        return_code = 0 if value.get("status") == "PASS" else 1
        profile_results[profile.name] = {
            "returnCode": return_code,
            "status": value.get("status"),
            "artifactClass": value.get("profile", {}).get("artifactClass"),
            "primaryFormat": value.get("profile", {}).get("primaryOutput", {}).get("format"),
        }
        profile_ok = profile_ok and value.get("status") == "PASS"
    record(checks, "delivery-profile-matrix", profile_ok, profiles=profile_results)

    request_example = ROOT / "artifact-delivery/examples/presentation-native-smoke-request-r1.json"
    with tempfile.TemporaryDirectory(prefix="artifact-delivery-request-doctor-") as tmp_text:
        tmp = Path(tmp_text)
        validation = provider.validate_delivery_request(request_example)
        compiled = provider.compile_delivery_plan(request_example)
        built = provider.execute_build_stage(request_example, tmp)
        details: dict[str, Any] = {
            "validateReturnCode": 0 if validation.get("status") == "PASS" else 1,
            "compileReturnCode": 0 if compiled.get("status") == "PASS" else 1,
            "buildReturnCode": 0 if built.get("status") == "PASS" else 1,
        }
        request_ok = all(
            value.get("status") == "PASS" for value in (validation, compiled, built)
        )
        if request_ok:
            try:
                artifact_path = Path(built["artifact"]["path"])
                details["artifact"] = built["artifact"]
                validator = selected_path(
                    "ARTIFACT_OPENXML_VALIDATOR", Path(OPENXML_LOCK["stableValidator"])
                )
                if Path(validator).exists():
                    validated = run([validator, str(artifact_path)])
                    details["openXmlReturnCode"] = validated.returncode
                    details["openXmlStdout"] = validated.stdout[-3000:]
                    request_ok = request_ok and validated.returncode == 0
                    verify_dir = tmp / "verify"
                    verify_value = provider.execute_verify_stage(
                        ROOT / "artifact-delivery/examples/pdu-sdu-presentation-r1.json",
                        artifact_path,
                        verify_dir,
                    )
                    details["verifyStageReturnCode"] = (
                        0 if verify_value.get("status") == "PASS" else 1
                    )
                    details["verifyStageStatus"] = verify_value.get("status")
                    details["verifyStagePending"] = verify_value.get(
                        "pendingRequiredGates"
                    )
                    required_local = {"profileSchema", "structural", "semantic"}
                    receipts = (
                        verify_value.get("receipts", {})
                        if isinstance(verify_value.get("receipts"), dict)
                        else {}
                    )
                    vsa_ok = (
                        verify_value.get("status") == "PASS"
                        and required_local.issubset(receipts)
                        and all(
                            receipts[name].get("verificationResult") == "PASSED"
                            for name in required_local
                        )
                        and all(
                            receipts[name].get("vsaValidation", {}).get("status")
                            == "PASS"
                            for name in required_local
                        )
                        and verify_value.get("profileVerificationComplete") is False
                    )
                    details["localVsaReceiptsPass"] = vsa_ok
                    request_ok = request_ok and vsa_ok
                else:
                    details["openXmlReturnCode"] = None
                    request_ok = False
            except Exception as error:
                details["parseError"] = str(error)
                request_ok = False
        else:
            details["validationFailures"] = validation.get("failures")
            details["compileFailures"] = compiled.get("failures")
            details["buildFailures"] = built.get("failures")
        record(
            checks,
            "digest-bound-request-presentation-source-build",
            request_ok,
            **details,
        )

    with tempfile.TemporaryDirectory(prefix="artifact-delivery-package-doctor-") as tmp_text:
        tmp = Path(tmp_text)
        base_profile = json.loads((ROOT / "artifact-delivery/examples/pdu-sdu-presentation-r1.json").read_text())
        base_profile.pop("$schema", None)
        base_profile["id"] = "presentation-package-doctor-r1"
        base_profile.pop("companions", None)
        for gate in list(base_profile["gates"]):
            base_profile["gates"][gate] = False
        for gate in ("profileSchema", "structural", "semantic"):
            base_profile["gates"][gate] = True
        package_profile = tmp / "profile.json"
        package_profile.write_text(json.dumps(base_profile))
        artifact_path = tmp / "artifact.pptx"
        built = provider.build_presentation_source(
            ROOT / "artifact-delivery/examples/presentation-native-smoke-source-r1.json",
            ROOT / "artifact-delivery/examples/pdu-sdu-presentation-r1.json",
            artifact_path,
        )
        verify_dir = tmp / "verify"
        verified = (
            provider.execute_verify_stage(
                package_profile, artifact_path, verify_dir
            )
            if built.get("status") == "PASS"
            else None
        )
        verify_report = tmp / "verify-stage.json"
        if verified is not None and verified.get("status") == "PASS":
            verify_report.write_text(
                json.dumps(verified, indent=2, sort_keys=True) + "\n"
            )
        package_dir = tmp / "package"
        packaged = (
            execute_oci_package_stage(
                package_profile,
                artifact_path,
                verify_report,
                package_dir,
                allow_local_unsigned=True,
            )
            if verify_report.is_file()
            else None
        )
        details = {
            "buildReturnCode": 0 if built.get("status") == "PASS" else 1,
            "verifyReturnCode": (
                0
                if verified is not None and verified.get("status") == "PASS"
                else None if verified is None else 1
            ),
            "packageReturnCode": (
                0
                if packaged is not None and packaged.get("status") == "PASS"
                else None if packaged is None else 1
            ),
        }
        package_ok = False
        if packaged is not None and packaged.get("status") == "PASS":
            try:
                package_value = packaged
                layout = package_dir / "layout"
                refs = package_value.get("oci", {}).get("discover", {}).get(
                    "referrers", []
                )
                package_ok = (
                    package_value.get("status") == "PASS"
                    and package_value.get("releaseReady") is False
                    and package_value.get("trustStanding")
                    == "LOCAL_UNSIGNED_DEVELOPMENT"
                    and (layout / "index.json").is_file()
                    and (layout / "oci-layout").is_file()
                    and not (package_dir / "package-index.json").exists()
                    and not (package_dir / "release-manifest.json").exists()
                    and len(refs) == 4
                )
                details["packageTrustStanding"] = package_value.get(
                    "trustStanding"
                )
                details["packageReleaseReady"] = package_value.get("releaseReady")
                details["ociSubjectDigest"] = (
                    package_value.get("oci", {}).get("subject", {}).get("digest")
                )
                details["ociReferrerCount"] = len(refs)
                details["legacyPackageIndexPresent"] = (
                    package_dir / "package-index.json"
                ).exists()
                details["legacyReleaseManifestPresent"] = (
                    package_dir / "release-manifest.json"
                ).exists()
            except Exception as error:
                details["parseError"] = str(error)
        else:
            details["packageFailures"] = (
                packaged.get("failures") if isinstance(packaged, dict) else None
            )
        record(
            checks,
            "digest-bound-development-oci-package",
            package_ok,
            **details,
        )

    pandoc = selected_path("ARTIFACT_PANDOC", GLOBAL_PANDOC, ROOT / ".cache/artifact-toolchain/pandoc/current/bin/pandoc")
    proc = run([pandoc, "--version"])
    first = proc.stdout.splitlines()[0] if proc.stdout else ""
    record(checks, "pandoc", proc.returncode == 0 and first == f"pandoc {LOCK['pandoc']['version']}", observed=first)

    verapdf = selected_path("ARTIFACT_VERAPDF", GLOBAL_VERAPDF, ROOT / ".cache/artifact-toolchain/verapdf/current/verapdf")
    proc = run([verapdf, "--version"])
    first = proc.stdout.splitlines()[0] if proc.stdout else ""
    record(checks, "verapdf", proc.returncode == 0 and first == f"veraPDF {LOCK['veraPDF']['version']}", observed=first)

    qpdf = os.environ.get("ARTIFACT_QPDF") or shutil.which("qpdf")
    if qpdf:
        proc = run([qpdf, "--version"])
        first = proc.stdout.splitlines()[0] if proc.stdout else ""
        record(checks, "qpdf", proc.returncode == 0 and LOCK["pdf"]["qpdf"] in first, observed=first)
    else:
        record(checks, "qpdf", False, error="qpdf not found")

    vnu = Path(selected_path("ARTIFACT_VNU", GLOBAL_VNU, ROOT / ".cache/artifact-toolchain/vnu/vnu.jar"))
    java = shutil.which("java")
    if vnu.is_file() and java:
        proc = run([java, "-jar", str(vnu), "--version"])
        observed_vnu = (proc.stdout.strip() or proc.stderr.strip()).splitlines()[0] if (proc.stdout.strip() or proc.stderr.strip()) else ""
        observed_digest = __import__("hashlib").sha256(vnu.read_bytes()).hexdigest()
        expected_vnu = LOCK["nuHtmlChecker"]
        ok = (
            proc.returncode == 0
            and observed_vnu == expected_vnu["version"]
            and observed_digest == expected_vnu["jarSha256"]
        )
        record(checks, "nu-html-checker", ok, observedVersion=observed_vnu, observedSha256=observed_digest, expected=expected_vnu)
    else:
        record(checks, "nu-html-checker", False, error="Nu Html Checker or Java not found")

    cosign = selected_path("ARTIFACT_COSIGN", GLOBAL_COSIGN, ROOT / ".cache/artifact-toolchain/cosign/current/bin/cosign")
    if not Path(cosign).is_file():
        cosign = shutil.which("cosign")
    if cosign:
        proc = run([cosign, "version"])
        output = proc.stdout + "\n" + proc.stderr
        match = re.search(r"GitVersion:\s*v(\d+)\.(\d+)\.(\d+)(?:[^\s]*)?", output)
        version_tuple = tuple(int(item) for item in match.groups()) if match else ()
        observed_version = ".".join(str(item) for item in version_tuple) if version_tuple else None
        observed_digest = hashlib.sha256(Path(cosign).read_bytes()).hexdigest()
        expected_cosign = LOCK["cosign"]
        standard_floor = tuple(int(item) for item in expected_cosign["standardizedBundleMinimumVersion"].split("."))
        source_text = (ROOT / "artifact_trust/vsa.py").read_text()
        standard_bundle_policy_ok = (
            proc.returncode == 0
            and bool(version_tuple)
            and version_tuple >= standard_floor
            and observed_version == expected_cosign["observedVersion"]
            and observed_digest == expected_cosign["binarySha256"]
            and expected_cosign.get("legacyBundlePolicy") == "REJECT"
            and "media_type != SIGSTORE_BUNDLE_V03" in source_text
            and '"--check-claims=true"' in source_text
            and "COSIGN_STANDARD_BUNDLE_MIN_VERSION = (3, 0, 6)" in source_text
        )
        record(
            checks,
            "cosign-standard-bundle-vsa-verifier",
            standard_bundle_policy_ok,
            observedVersion=observed_version,
            observedSha256=observed_digest,
            minimum=".".join(str(item) for item in standard_floor),
            legacyBundlePolicy=expected_cosign.get("legacyBundlePolicy"),
            publicKeyStanding=expected_cosign.get("publicKeyStandardBundleStanding"),
            keylessStanding=expected_cosign.get("keylessStandardBundleStanding"),
        )
        runtime_cosign = artifact_trust_vsa.cosign_tool_fact()
        provenance = runtime_cosign.get("provenance", {}) if isinstance(runtime_cosign, dict) else {}
        provenance_ok = runtime_cosign.get("status") == "PASS" and provenance.get("status") == "PASS"
        record(
            checks,
            "cosign-selection-provenance",
            provenance_ok,
            selectedVerifier=runtime_cosign.get("path"),
            selectedVerifierSha256=runtime_cosign.get("sha256"),
            provenance=provenance,
        )
        keyless_policy_ok = (
            standard_bundle_policy_ok
            and expected_cosign.get("keylessStandardBundleStanding") == "NOT_EXERCISED_IDENTITY_ISSUER_TLOG_REQUIRED"
            and "certificateIdentity" in source_text
            and "certificateOidcIssuer" in source_text
            and "requireTransparencyLog" in source_text
        )
        record(
            checks,
            "cosign-keyless-standard-bundle-policy",
            keyless_policy_ok,
            observedVersion=observed_version,
            legacyKeylessPatchedVersion=expected_cosign.get("legacyKeylessPatchedVersion"),
            standing=expected_cosign.get("keylessStandardBundleStanding"),
            note="Legacy bundles remain rejected; standardized-bundle keyless verification is not claimed until identity/issuer/tlog integration is exercised.",
        )
    else:
        record(checks, "cosign-standard-bundle-vsa-verifier", False, error="cosign not found")
        record(checks, "cosign-keyless-standard-bundle-policy", False, error="cosign not found")

    signing_config = ROOT / "artifact-delivery/sigstore-local-signing-config-v1.json"
    if signing_config.is_file():
        signing_value = json.loads(signing_config.read_text())
        signing_digest = hashlib.sha256(signing_config.read_bytes()).hexdigest()
        service_empty = all(signing_value.get(field) == [] for field in ("caUrls", "oidcUrls", "rekorTlogUrls", "tsaUrls"))
        signing_ok = (
            signing_value.get("mediaType") == "application/vnd.dev.sigstore.signingconfig.v0.2+json"
            and service_empty
            and signing_digest == LOCK["attestation"]["localSigningConfigSha256"]
        )
        record(
            checks,
            "sigstore-local-signing-config",
            signing_ok,
            sha256=signing_digest,
            externalServicesEmpty=service_empty,
        )
    else:
        record(checks, "sigstore-local-signing-config", False, error="local Sigstore signing config is absent")

    dotnet = selected_path("ARTIFACT_DOTNET", Path(OPENXML_LOCK["stableDotnet"]))
    proc = run([dotnet, "--version"])
    observed_dotnet = proc.stdout.strip()
    dotnet_ok = proc.returncode == 0 and observed_dotnet.startswith(OPENXML_LOCK["dotnetPolicy"]["sdkMajorMinor"] + ".")
    record(checks, "dotnet", dotnet_ok, observed=observed_dotnet, policy=OPENXML_LOCK["dotnetPolicy"])
    openxml_validator = selected_path("ARTIFACT_OPENXML_VALIDATOR", Path(OPENXML_LOCK["stableValidator"]))
    record(checks, "openxml-validator-carrier", Path(openxml_validator).is_file(), selected=openxml_validator, methodId=OPENXML_LOCK["methodId"])

    with tempfile.TemporaryDirectory(prefix="artifact-toolchain-doctor-") as tmp_text:
        tmp = Path(tmp_text)
        pptx = tmp / "probe.pptx"
        xlsx = tmp / "probe.xlsx"
        md = tmp / "probe.md"
        docx = tmp / "probe.docx"
        md.write_text("# Artifact E2E\n\nPandoc DOCX composition smoke.\n")
        generator = """
from pptx import Presentation
from pptx.util import Inches
import xlsxwriter, sys
pptx_path, xlsx_path = sys.argv[1:3]
prs = Presentation(); prs.slide_width = Inches(13.333333); prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6]); slide.shapes.add_textbox(Inches(.5), Inches(.5), Inches(6), Inches(.5)).text = 'Artifact E2E'
prs.save(pptx_path)
wb = xlsxwriter.Workbook(xlsx_path); ws = wb.add_worksheet(); fmt = wb.add_format({'bold': True, 'font_name': 'Arial'}); ws.write('A1', 'Artifact E2E', fmt); wb.close()
"""
        generated = run([python, "-c", generator, str(pptx), str(xlsx)])
        converted = run([pandoc, str(md), "-o", str(docx)])
        validations: dict[str, dict[str, Any]] = {}
        if generated.returncode == 0 and converted.returncode == 0 and dotnet_ok and Path(openxml_validator).is_file():
            for artifact in (pptx, docx, xlsx):
                validated = run([openxml_validator, str(artifact)])
                validations[artifact.suffix] = {
                    "returnCode": validated.returncode,
                    "stdout": validated.stdout[-3000:],
                    "stderr": validated.stderr[-1000:],
                }
        cross_ok = (
            generated.returncode == 0
            and converted.returncode == 0
            and len(validations) == 3
            and all(item["returnCode"] == 0 for item in validations.values())
        )
        record(
            checks,
            "cross-format-ooxml-generation-validation",
            cross_ok,
            generatedReturnCode=generated.returncode,
            pandocReturnCode=converted.returncode,
            validations=validations,
        )

    node = shutil.which("node")
    if node:
        node_env = os.environ.copy()
        if "ARTIFACT_NODE_PACKAGE_ROOT" not in node_env and (GLOBAL_NODE_PACKAGE_ROOT / "package.json").is_file():
            node_env["ARTIFACT_NODE_PACKAGE_ROOT"] = str(GLOBAL_NODE_PACKAGE_ROOT)
        proc = run([node, "./probe.mjs"], cwd=ROOT / "artifact-delivery/node", timeout=90, env=node_env)
        try:
            web = json.loads(proc.stdout)
        except Exception:
            web = {"ok": False, "stdout": proc.stdout[-3000:], "stderr": proc.stderr[-2000:]}
        chromium = web.get("browsers", {}).get("chromium", {})
        firefox = web.get("browsers", {}).get("firefox", {})
        webkit = web.get("browsers", {}).get("webkit", {})
        ok = (
            proc.returncode == 0
            and web.get("ok") is True
            and chromium.get("status") == "PASS"
            and chromium.get("accessibility", {}).get("violationCount") == 0
            and firefox.get("status") == "PASS"
            and webkit.get("status") in {"LOCAL_HOST_COMPATIBILITY_NOT_PROVEN", "NOT_INSTALLED"}
        )
        record(checks, "web-browser-matrix", ok, browsers=web.get("browsers", {}), stderr=proc.stderr[-1000:])
    else:
        record(checks, "web-browser-matrix", False, error="node not found")

    web_request = ROOT / "artifact-delivery/examples/web-smoke-request-r1.json"
    with tempfile.TemporaryDirectory(
        prefix="artifact-delivery-web-profile-doctor-"
    ) as tmp_text:
        tmp = Path(tmp_text)
        build_dir = tmp / "build"
        verify_dir = tmp / "verify"
        built = provider.execute_build_stage(web_request, build_dir)
        details: dict[str, Any] = {
            "buildReturnCode": 0 if built.get("status") == "PASS" else 1
        }
        web_profile_ok = built.get("status") == "PASS"
        if web_profile_ok:
            try:
                artifact_path = Path(built["artifact"]["path"])
                verify_value = provider.execute_verify_stage(
                    ROOT / "artifact-delivery/examples/web-r1.json",
                    artifact_path,
                    verify_dir,
                )
                details["verifyReturnCode"] = (
                    0 if verify_value.get("status") == "PASS" else 1
                )
                receipts = verify_value.get("receipts", {})
                local_pass_gates = {
                    "profileSchema",
                    "structural",
                    "conformance",
                    "accessibility",
                }
                common_ok = (
                    local_pass_gates.issubset(receipts)
                    and all(
                        receipts[name].get("verificationResult") == "PASSED"
                        for name in local_pass_gates
                    )
                    and all(
                        receipts[name].get("vsaValidation", {}).get("status")
                        == "PASS"
                        for name in local_pass_gates
                    )
                )
                target = receipts.get("target", {})
                target_raw_path = Path(
                    target.get("rawEvidence", {}).get("path", "")
                )
                target_raw = (
                    json.loads(target_raw_path.read_text())
                    if target_raw_path.is_file()
                    else {}
                )
                chromium_target = target_raw.get("browserResults", {}).get(
                    "Chromium", {}
                )
                firefox_target = target_raw.get("browserResults", {}).get(
                    "Firefox", {}
                )
                webkit_target = target_raw.get("browserResults", {}).get(
                    "WebKit", {}
                )
                primary_ok = (
                    chromium_target.get("status") == "PASS"
                    and firefox_target.get("status") == "PASS"
                )
                target_full_pass = (
                    target.get("verificationResult") == "PASSED"
                    and webkit_target.get("status") == "PASS"
                )
                expected_local_block = (
                    target.get("verificationResult") == "FAILED"
                    and webkit_target.get("status")
                    in {
                        "FAIL",
                        "NOT_INSTALLED",
                        "LOCAL_HOST_COMPATIBILITY_NOT_PROVEN",
                    }
                    and verify_value.get("failures")
                    == ["executed verifier gate(s) failed: target"]
                )
                web_profile_ok = common_ok and primary_ok and (
                    target_full_pass or expected_local_block
                )
                details.update(
                    {
                        "verifyStatus": verify_value.get("status"),
                        "profileVerificationComplete": verify_value.get(
                            "profileVerificationComplete"
                        ),
                        "pendingRequiredGates": verify_value.get(
                            "pendingRequiredGates"
                        ),
                        "receipts": {
                            name: item.get("verificationResult")
                            for name, item in sorted(receipts.items())
                        },
                        "targetBrowsers": {
                            "Chromium": chromium_target.get("status"),
                            "Firefox": firefox_target.get("status"),
                            "WebKit": webkit_target.get("status"),
                        },
                        "acceptedBoundary": (
                            "FULL_TARGET_PASS"
                            if target_full_pass
                            else "WEBKIT_SUPPORTED_RUNNER_REQUIRED"
                            if expected_local_block
                            else "UNEXPECTED"
                        ),
                    }
                )
            except Exception as error:
                details["parseError"] = str(error)
                web_profile_ok = False
        else:
            details["buildFailures"] = built.get("failures")
        record(
            checks,
            "web-profile-local-verification-boundary",
            web_profile_ok,
            **details,
        )

    status = "PASS" if all(item["status"] == "PASS" for item in checks) else "FAIL"
    result = {
        "schemaVersion": 1,
        "kind": "artifact-build-delivery-toolchain-doctor",
        "status": status,
        "checks": checks,
        "boundary": "Mechanical local tool/profile/primary-runner evidence only; Office target, visual acceptance, PDF profile conformance, WebKit supported-runner acceptance and destination delivery remain separate gates.",
    }
    print(json.dumps(result, indent=2))
    return 0 if status == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())
