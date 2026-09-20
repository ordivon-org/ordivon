from __future__ import annotations

import hashlib
import re
import zipfile
from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path, PurePosixPath
from typing import Any

from artifact_core.contracts import file_fact, sha256_file

ROOT = Path(__file__).resolve().parents[2]
DEFAULT_PRESENTATION_SOURCE_SCHEMA = ROOT / "artifact-delivery/presentation-source-v1.schema.json"
DEFAULT_PRESENTATION_SEMANTIC_SVG_SOURCE_SCHEMA = ROOT / "artifact-delivery/presentation-semantic-svg-source-v1.schema.json"
PPT_MASTER_PROVIDER_LOCK = ROOT / "artifact-delivery/ppt-master-provider-v1.lock.json"


@dataclass(frozen=True)
class PresentationBuildHooks:
    validate_json_document: Callable[..., dict[str, Any]]
    inspect_pptx: Callable[..., dict[str, Any]]
    verify_semantics: Callable[[dict[str, Any], dict[str, Any]], dict[str, Any]]
    normalize_python_pptx: Callable[[Path], dict[str, Any]]
    canonicalize_ppt_master: Callable[[Path], dict[str, Any]]

def _resolve_semantic_svg_source_path(source_path: Path, relative: str) -> Path:
    candidate = Path(relative)
    if candidate.is_absolute():
        raise RuntimeError(f"semantic SVG source path must be relative: {relative!r}")
    source_root = source_path.resolve().parent
    resolved = (source_root / candidate).resolve()
    try:
        resolved.relative_to(source_root)
    except ValueError as error:
        raise RuntimeError(f"semantic SVG source path escapes source directory: {relative!r}") from error
    return resolved


def _safe_project_relative_path(value: str) -> PurePosixPath:
    path = PurePosixPath(value)
    if path.is_absolute() or not path.parts or any(part in {"", ".", ".."} for part in path.parts):
        raise RuntimeError(f"unsafe provider project-relative path: {value!r}")
    return path


def _semantic_svg_source_material_facts(source_path: Path, source: dict[str, Any]) -> tuple[list[dict[str, Any]], list[str]]:
    materials: list[dict[str, Any]] = []
    failures: list[str] = []
    seen_page_ids: set[str] = set()
    for index, page in enumerate(source.get("pages", []) if isinstance(source.get("pages"), list) else []):
        page_id = str(page.get("id", ""))
        if page_id in seen_page_ids:
            failures.append(f"duplicate semantic SVG page id: {page_id}")
        seen_page_ids.add(page_id)
        try:
            page_path = _resolve_semantic_svg_source_path(source_path, str(page.get("path", "")))
            if page_path.suffix.lower() != ".svg":
                failures.append(f"semantic SVG page {page_id or index + 1} must use .svg")
            actual = sha256_file(page_path)
            if actual != page.get("sha256"):
                failures.append(f"semantic SVG page digest mismatch: {page_id or index + 1}")
            fact = file_fact(page_path)
            fact["purpose"] = "semantic-svg-page"
            fact["pageId"] = page_id
            materials.append(fact)
        except Exception as error:
            failures.append(f"semantic SVG page reference error ({page_id or index + 1}): {error}")
    seen_targets: set[str] = set()
    for item in source.get("materials", []) if isinstance(source.get("materials"), list) else []:
        try:
            target = _safe_project_relative_path(str(item.get("projectRelativePath", "")))
            target_key = target.as_posix()
            if target_key.startswith("svg_output/"):
                failures.append(f"semantic SVG provider material may not target reserved svg_output/: {target_key}")
            if target_key in seen_targets:
                failures.append(f"duplicate semantic SVG provider material target: {target_key}")
            seen_targets.add(target_key)
            material_path = _resolve_semantic_svg_source_path(source_path, str(item.get("path", "")))
            actual = sha256_file(material_path)
            if actual != item.get("sha256"):
                failures.append(f"semantic SVG provider material digest mismatch: {item.get('path')}")
            fact = file_fact(material_path)
            fact["purpose"] = item.get("purpose") or "semantic-svg-material"
            fact["projectRelativePath"] = target_key
            materials.append(fact)
        except Exception as error:
            failures.append(f"semantic SVG provider material reference error: {error}")
    return materials, failures


def _resolve_presentation_template(source_path: Path, template: dict[str, Any]) -> tuple[Path | None, list[str]]:
    failures: list[str] = []
    if template.get("format") != "pptx":
        failures.append("presentation template format must be pptx for the current python-pptx adapter")
    path_value = template.get("path")
    digest_value = template.get("sha256")
    if not isinstance(path_value, str) or not path_value:
        failures.append("presentation template path is required")
        return None, failures
    path = Path(path_value)
    if not path.is_absolute():
        path = source_path.resolve().parent / path
    path = path.resolve()
    if path.suffix.lower() != ".pptx":
        failures.append("current presentation template adapter accepts a slide-free .pptx authority; true .potx ingestion is not claimed")
    if not path.is_file():
        failures.append(f"presentation template is absent: {path}")
        return path, failures
    actual = sha256_file(path)
    if not isinstance(digest_value, str) or actual != digest_value:
        failures.append(f"presentation template digest mismatch: expected {digest_value}, got {actual}")
    return path, failures


def _resolve_presentation_image(source_path: Path, element: dict[str, Any]) -> tuple[Path | None, list[str]]:
    failures: list[str] = []
    path_value = element.get("path")
    digest_value = element.get("sha256")
    if not isinstance(path_value, str) or not path_value:
        return None, ["presentation image path is required"]
    path = Path(path_value)
    if not path.is_absolute():
        path = source_path.resolve().parent / path
    path = path.resolve()
    if path.suffix.lower() not in {".png", ".jpg", ".jpeg"}:
        failures.append(f"presentation image format is not admitted by v1 raster adapter: {path.suffix.lower()}")
    if not path.is_file():
        failures.append(f"presentation image is absent: {path}")
        return path, failures
    actual = sha256_file(path)
    if not isinstance(digest_value, str) or actual != digest_value:
        failures.append(f"presentation image digest mismatch: expected {digest_value}, got {actual}")
    try:
        from PIL import Image
        with Image.open(path) as image:
            image.verify()
    except Exception as error:
        failures.append(f"presentation image bytes are not a valid admitted raster image: {error}")
    return path, failures


def _presentation_source_material_facts(source_path: Path, source: dict[str, Any]) -> tuple[list[dict[str, Any]], list[str]]:
    materials: list[dict[str, Any]] = []
    failures: list[str] = []
    template = source.get("template")
    if isinstance(template, dict):
        template_path, template_failures = _resolve_presentation_template(source_path, template)
        failures.extend(template_failures)
        if template_path is not None and template_path.is_file() and not template_failures:
            materials.append(file_fact(template_path))
    for slide in source.get("slides", []) if isinstance(source.get("slides"), list) else []:
        for element in slide.get("elements", []) if isinstance(slide.get("elements"), list) else []:
            if element.get("kind") != "image":
                continue
            image_path, image_failures = _resolve_presentation_image(source_path, element)
            failures.extend(f"{slide.get('id')}/{element.get('id')}: {item}" for item in image_failures)
            if image_path is not None and image_path.is_file() and not image_failures:
                fact = file_fact(image_path)
                if not any(existing.get("path") == fact.get("path") and existing.get("digest") == fact.get("digest") for existing in materials):
                    materials.append(fact)
    return materials, failures


def _presentation_template_package_fact(path: Path) -> dict[str, Any]:
    theme_parts: list[dict[str, Any]] = []
    master_parts: list[dict[str, Any]] = []
    layout_parts: list[dict[str, Any]] = []
    with zipfile.ZipFile(path) as package:
        for name, target in (
            (r"ppt/theme/theme\d+\.xml", theme_parts),
            (r"ppt/slideMasters/slideMaster\d+\.xml", master_parts),
            (r"ppt/slideLayouts/slideLayout\d+\.xml", layout_parts),
        ):
            pattern = re.compile(name)
            for member in sorted(item for item in package.namelist() if pattern.fullmatch(item)):
                data = package.read(member)
                target.append({"part": member, "sha256": hashlib.sha256(data).hexdigest(), "size": len(data)})
    return {
        "artifact": file_fact(path),
        "themeParts": theme_parts,
        "masterParts": master_parts,
        "layoutParts": layout_parts,
    }


def _presentation_layout_by_id(prs: Any, layout_id: str) -> tuple[Any | None, list[str]]:
    wanted = layout_id.casefold()
    matches = [layout for layout in prs.slide_layouts if str(layout.name).casefold() == wanted]
    if len(matches) == 1:
        return matches[0], []
    if not matches:
        return None, [f"presentation layout not found in selected template/master set: {layout_id}"]
    return None, [f"presentation layout name is ambiguous in selected template/master set: {layout_id}"]


def _presentation_layout_fact(layout: Any) -> dict[str, Any]:
    master = layout.slide_master
    placeholders = []
    for placeholder in layout.placeholders:
        placeholders.append({
            "idx": int(placeholder.placeholder_format.idx),
            "name": str(placeholder.name),
            "type": str(placeholder.placeholder_format.type),
        })
    return {
        "name": str(layout.name),
        "part": str(layout.part.partname),
        "masterPart": str(master.part.partname),
        "placeholders": placeholders,
    }


def _hex_color(value: str):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(value.upper())


def _apply_text_to_shape(shape: Any, element: dict[str, Any], align_map: dict[str, Any]) -> None:
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Pt

    frame = shape.text_frame
    frame.clear()
    paragraph = frame.paragraphs[0]
    paragraph.text = element.get("text", "")
    if element.get("align") is not None:
        paragraph.alignment = align_map.get(element.get("align"), PP_ALIGN.LEFT)
    runs = list(paragraph.runs)
    if not runs:
        runs = [paragraph.add_run()]
    for run in runs:
        if element.get("fontFamily") is not None:
            run.font.name = element["fontFamily"]
        if element.get("fontSizePt") is not None:
            run.font.size = Pt(float(element["fontSizePt"]))
        if "bold" in element:
            run.font.bold = bool(element["bold"])
        if "italic" in element:
            run.font.italic = bool(element["italic"])
        if element.get("colorHex"):
            run.font.color.rgb = _hex_color(str(element["colorHex"]))
        if element.get("opacity") is not None:
            opacity = float(element["opacity"])
            if opacity < 0 or opacity > 1:
                raise ValueError(f"presentation text opacity must be between 0 and 1: {opacity}")
            # DrawingML alpha is the standard 0..100000 opacity scalar. python-pptx
            # creates a distinct run after each soft line break, so every run must
            # carry the same resolved font/style/alpha rather than styling only the
            # first run and letting later lines fall back to visible defaults.
            if opacity < 1:
                from pptx.oxml.ns import qn
                from pptx.oxml.xmlchemy import OxmlElement
                r_pr = run._r.get_or_add_rPr()
                solid_fill = r_pr.find(qn("a:solidFill"))
                if solid_fill is None or len(solid_fill) == 0:
                    raise RuntimeError("text opacity requires a resolved solid font color")
                color_node = solid_fill[0]
                for existing in list(color_node.findall(qn("a:alpha"))):
                    color_node.remove(existing)
                alpha = OxmlElement("a:alpha")
                alpha.set("val", str(int(round(opacity * 100000))))
                color_node.append(alpha)

