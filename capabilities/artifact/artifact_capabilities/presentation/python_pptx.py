from __future__ import annotations

import importlib.metadata
from pathlib import Path
from typing import Any

from artifact_core.contracts import file_fact
from artifact_core.profile_v1 import validate_profile_v1 as validate_profile
from .common import (
    DEFAULT_PRESENTATION_SOURCE_SCHEMA,
    PresentationBuildHooks,
    _apply_text_to_shape,
    _presentation_layout_by_id,
    _presentation_layout_fact,
    _presentation_template_package_fact,
    _resolve_presentation_image,
    _resolve_presentation_template,
)

def build_presentation_source(
    source_path: Path,
    profile_path: Path,
    output_path: Path,
    source_schema_path: Path = DEFAULT_PRESENTATION_SOURCE_SCHEMA,
    *,
    hooks: PresentationBuildHooks,
) -> dict[str, Any]:
    source_result = hooks.validate_json_document(source_path, source_schema_path, "presentation-source")
    profile_result = validate_profile(profile_path)
    failures: list[str] = []
    source = source_result.get("document", {})
    profile = profile_result.get("profile", {})
    if source_result.get("status") != "PASS":
        failures.append("presentation source schema did not PASS")
    if profile_result.get("status") != "PASS":
        failures.append("delivery profile schema did not PASS")
    if source.get("sourceMode") != "native-composition":
        failures.append("v1 native builder accepts sourceMode=native-composition only")
    if source.get("profileId") != profile.get("id"):
        failures.append("presentation source profileId does not match selected profile")
    if source.get("aspectRatio") != profile.get("aspectRatio"):
        failures.append("presentation source aspectRatio does not match selected profile")
    declared_fonts = {str(item.get("family")) for item in profile.get("fonts", [])}
    width = float(source.get("slideSizeInches", {}).get("width", 0))
    height = float(source.get("slideSizeInches", {}).get("height", 0))
    if width <= 0 or height <= 0:
        failures.append("presentation slide size must be positive")
    template_path: Path | None = None
    template_spec = source.get("template")
    if isinstance(template_spec, dict):
        template_path, template_failures = _resolve_presentation_template(source_path, template_spec)
        failures.extend(template_failures)
    slide_ids: set[str] = set()
    for slide in source.get("slides", []) if isinstance(source.get("slides"), list) else []:
        slide_id = str(slide.get("id"))
        if slide_id in slide_ids:
            failures.append(f"duplicate slide id: {slide_id}")
        slide_ids.add(slide_id)
        if template_path is not None and not slide.get("layoutId"):
            failures.append(f"template-bound slide requires layoutId: {slide_id}")
        element_ids: set[str] = set()
        for element in slide.get("elements", []) if isinstance(slide.get("elements"), list) else []:
            element_id = str(element.get("id"))
            if element_id in element_ids:
                failures.append(f"duplicate element id on {slide_id}: {element_id}")
            element_ids.add(element_id)
            kind = element.get("kind")
            if kind == "text":
                family = element.get("fontFamily")
                if family is not None and str(family) not in declared_fonts:
                    failures.append(f"undeclared font family on {slide_id}/{element_id}: {family}")
                if element.get("opacity") is not None and element.get("colorHex") is None:
                    failures.append(f"text opacity requires explicit colorHex on {slide_id}/{element_id}")
                needs_box = element.get("placeholderIdx") is None
            elif kind == "image":
                _, image_failures = _resolve_presentation_image(source_path, element)
                failures.extend(f"{slide_id}/{element_id}: {item}" for item in image_failures)
                needs_box = True
            else:
                failures.append(f"unsupported presentation element kind on {slide_id}/{element_id}: {kind}")
                needs_box = False
            if needs_box:
                box = element.get("box", {})
                x, y, w, h = (float(box.get(key, 0)) for key in ("x", "y", "w", "h"))
                if x < 0 or y < 0 or w <= 0 or h <= 0 or x + w > width + 1e-6 or y + h > height + 1e-6:
                    failures.append(f"out-of-bounds box on {slide_id}/{element_id}")
    if failures:
        return {
            "status": "FAIL",
            "source": file_fact(source_path),
            "profile": file_fact(profile_path),
            "failures": failures,
        }
    from pptx import Presentation
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches

    align_map = {
        "left": PP_ALIGN.LEFT,
        "center": PP_ALIGN.CENTER,
        "right": PP_ALIGN.RIGHT,
        "justify": PP_ALIGN.JUSTIFY,
    }
    prs = Presentation(str(template_path)) if template_path is not None else Presentation()
    template_fact: dict[str, Any] | None = None
    if template_path is not None:
        if len(prs.slides) != 0:
            failures.append("presentation template authority must contain zero slides in v1 to prevent sample-slide contamination")
        template_width = float(prs.slide_width) / 914400.0
        template_height = float(prs.slide_height) / 914400.0
        if abs(template_width - width) > 1e-5 or abs(template_height - height) > 1e-5:
            failures.append(
                f"presentation template slide size mismatch: template={template_width:.6f}x{template_height:.6f}, source={width:.6f}x{height:.6f}"
            )
        template_fact = _presentation_template_package_fact(template_path)
    else:
        prs.slide_width = Inches(width)
        prs.slide_height = Inches(height)
    layout_bindings: list[dict[str, Any]] = []
    planned: list[tuple[dict[str, Any], Any]] = []
    for slide_spec in source.get("slides", []):
        layout_id = str(slide_spec.get("layoutId") or "Blank")
        layout, layout_failures = _presentation_layout_by_id(prs, layout_id)
        failures.extend(f"{slide_spec.get('id')}: {item}" for item in layout_failures)
        if layout is None:
            continue
        placeholder_indices = {int(p.placeholder_format.idx): p for p in layout.placeholders}
        for element in slide_spec.get("elements", []):
            if element.get("kind") == "text" and element.get("placeholderIdx") is not None:
                idx = int(element["placeholderIdx"])
                placeholder = placeholder_indices.get(idx)
                if placeholder is None:
                    failures.append(f"{slide_spec.get('id')}/{element.get('id')}: placeholder idx {idx} is absent from layout {layout.name}")
                elif not placeholder.has_text_frame:
                    failures.append(f"{slide_spec.get('id')}/{element.get('id')}: placeholder idx {idx} on layout {layout.name} has no text frame")
        planned.append((slide_spec, layout))
        layout_bindings.append({"slideId": slide_spec.get("id"), "layout": _presentation_layout_fact(layout)})
    if failures:
        return {
            "status": "FAIL",
            "source": file_fact(source_path),
            "profile": file_fact(profile_path),
            "template": template_fact,
            "layoutBindings": layout_bindings,
            "failures": failures,
        }
    media_bindings: list[dict[str, Any]] = []
    for slide_spec, layout in planned:
        slide = prs.slides.add_slide(layout)
        placeholders = {int(p.placeholder_format.idx): p for p in slide.placeholders}
        for element in slide_spec.get("elements", []):
            kind = element.get("kind")
            if kind == "text":
                if element.get("placeholderIdx") is not None:
                    shape = placeholders[int(element["placeholderIdx"])]
                    _apply_text_to_shape(shape, element, align_map)
                else:
                    box = element["box"]
                    shape = slide.shapes.add_textbox(Inches(box["x"]), Inches(box["y"]), Inches(box["w"]), Inches(box["h"]))
                    _apply_text_to_shape(shape, element, align_map)
            elif kind == "image":
                image_path, image_failures = _resolve_presentation_image(source_path, element)
                if image_failures or image_path is None:
                    raise RuntimeError(f"presentation image binding changed after validation: {image_failures}")
                box = element["box"]
                shape = slide.shapes.add_picture(str(image_path), Inches(box["x"]), Inches(box["y"]), Inches(box["w"]), Inches(box["h"]))
                # python-pptx exposes native descr metadata but does not provide a stable high-level
                # setter for decorative state. Preserve user-authored accessibility intent in the
                # receipt and use descr for non-decorative images; accessibility remains an
                # independent target gate rather than being inferred from this field alone.
                if not bool(element.get("decorative")) and element.get("altText"):
                    cNvPr = shape._element.xpath('.//p:cNvPr')[0]
                    cNvPr.set('descr', str(element["altText"]))
                media_bindings.append({
                    "slideId": slide_spec.get("id"),
                    "elementId": element.get("id"),
                    "artifact": file_fact(image_path),
                    "decorative": bool(element.get("decorative")),
                    "altText": element.get("altText"),
                    "shapeId": int(shape.shape_id),
                })
            else:
                raise RuntimeError(f"unsupported presentation element kind: {kind}")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)
    container_normalization = hooks.normalize_python_pptx(output_path)
    built = hooks.inspect_pptx(output_path, profile.get("semanticPolicy", {}).get("placeholderPatterns", []))
    semantic = hooks.verify_semantics(profile, built)
    post_failures: list[str] = []
    if built.get("status") != "PASS":
        post_failures.append("built PPTX failed package/relationship inspection")
    if semantic.get("status") != "PASS":
        post_failures.append("built PPTX failed presentation semantic checks")
    return {
        "status": "PASS" if not post_failures else "FAIL",
        "source": file_fact(source_path),
        "profile": file_fact(profile_path),
        "artifact": file_fact(output_path),
        "presentationId": source.get("presentationId"),
        "slideCount": len(source.get("slides", [])),
        "builder": {"implementation": "python-pptx", "version": importlib.metadata.version("python-pptx")},
        "containerNormalization": container_normalization,
        "template": template_fact,
        "layoutBindings": layout_bindings,
        "mediaBindings": media_bindings,
        "inspection": built,
        "semantic": semantic,
        "failures": post_failures,
        "boundary": "Builder PASS establishes source/profile binding plus native PPTX package/semantic checks. Open XML SDK, target PowerPoint, visual, accessibility and delivery gates remain independent.",
    }

