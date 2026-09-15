from __future__ import annotations

import importlib.util
import json
import os
from pathlib import Path
import shutil
import subprocess
import tempfile
import unittest
import zipfile

ROOT = Path(__file__).resolve().parents[1]
MODULE_PATH = ROOT / "scripts/artifact_delivery.py"
SPEC = importlib.util.spec_from_file_location("artifact_delivery", MODULE_PATH)
assert SPEC is not None and SPEC.loader is not None
MODULE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(MODULE)


def openxml_validator_path() -> Path:
    return Path(os.environ.get("ARTIFACT_OPENXML_VALIDATOR", "/root/.local/share/ordivon-workstation/artifact-openxml-v1/current/bin/validate-openxml"))


def openxml_dotnet_path() -> Path:
    return Path(os.environ.get("ARTIFACT_DOTNET", "/root/.local/share/ordivon-workstation/artifact-openxml-v1/current/bin/dotnet"))


PRESENTATION_XML = """<?xml version='1.0' encoding='UTF-8' standalone='yes'?>
<p:presentation xmlns:a='http://schemas.openxmlformats.org/drawingml/2006/main'
 xmlns:r='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
 xmlns:p='http://schemas.openxmlformats.org/presentationml/2006/main'>
 <p:sldIdLst><p:sldId id='256' r:id='rId1'/></p:sldIdLst>
 <p:sldSz cx='12192000' cy='6858000'/>
</p:presentation>
"""
SLIDE_XML = """<?xml version='1.0' encoding='UTF-8' standalone='yes'?>
<p:sld xmlns:a='http://schemas.openxmlformats.org/drawingml/2006/main'
 xmlns:r='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
 xmlns:p='http://schemas.openxmlformats.org/presentationml/2006/main'>
 <p:cSld><p:spTree><p:sp><p:txBody><a:p><a:r><a:rPr typeface='Microsoft YaHei'/><a:t>正式内容</a:t></a:r></a:p></p:txBody></p:sp></p:spTree></p:cSld>
</p:sld>
"""
ROOT_RELS = """<?xml version='1.0' encoding='UTF-8' standalone='yes'?>
<Relationships xmlns='http://schemas.openxmlformats.org/package/2006/relationships'>
 <Relationship Id='rId1' Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument' Target='ppt/presentation.xml'/>
</Relationships>
"""
PRESENTATION_RELS = """<?xml version='1.0' encoding='UTF-8' standalone='yes'?>
<Relationships xmlns='http://schemas.openxmlformats.org/package/2006/relationships'>
 <Relationship Id='rId1' Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide' Target='slides/slide1.xml'/>
</Relationships>
"""
CONTENT_TYPES = """<?xml version='1.0' encoding='UTF-8' standalone='yes'?>
<Types xmlns='http://schemas.openxmlformats.org/package/2006/content-types'>
 <Default Extension='rels' ContentType='application/vnd.openxmlformats-package.relationships+xml'/>
 <Default Extension='xml' ContentType='application/xml'/>
</Types>
"""


def make_pptx(path: Path, text: str = "正式内容", font: str = "Microsoft YaHei") -> None:
    slide = SLIDE_XML.replace("正式内容", text).replace("Microsoft YaHei", font)
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("[Content_Types].xml", CONTENT_TYPES)
        z.writestr("_rels/.rels", ROOT_RELS)
        z.writestr("ppt/presentation.xml", PRESENTATION_XML)
        z.writestr("ppt/_rels/presentation.xml.rels", PRESENTATION_RELS)
        z.writestr("ppt/slides/slide1.xml", slide)


class ArtifactDeliveryTests(unittest.TestCase):
    def test_profile_declares_no_universal_document_model(self) -> None:
        profile = json.loads((ROOT / "artifact-delivery/examples/pdu-sdu-presentation-r1.json").read_text())
        self.assertEqual(profile["artifactClass"], "presentation")
        self.assertEqual(profile["primaryOutput"]["format"], "pptx")
        self.assertEqual(profile["targetRenderer"]["name"], "Microsoft PowerPoint Desktop")
        self.assertTrue(profile["gates"]["target"])
        self.assertTrue(profile["gates"]["deliveryReadback"])
        self.assertNotIn("documentAst", profile)
        self.assertNotIn("slideModel", profile)

    def test_cross_format_profiles_validate_under_shared_delivery_contract(self) -> None:
        profiles = {
            "pdu-sdu-presentation-r1.json": ("presentation", "pptx"),
            "document-r1.json": ("document", "docx"),
            "spreadsheet-r1.json": ("spreadsheet", "xlsx"),
            "web-r1.json": ("web", "html"),
            "pdf-fixed-r1.json": ("fixed-view", "pdf"),
            "pdf-accessible-r1.json": ("accessible", "pdf-ua-2"),
        }
        for filename, expected in profiles.items():
            with self.subTest(profile=filename):
                result = MODULE.validate_profile(ROOT / "artifact-delivery/examples" / filename)
                self.assertEqual(result["minimalContractErrors"], [], result)
                if result["jsonSchema"]["validator"] == "unavailable":
                    self.assertEqual(result["status"], "FAIL")
                    self.assertEqual(result["jsonSchema"]["status"], "NOT_RUN")
                else:
                    self.assertEqual(result["status"], "PASS", result)
                    self.assertEqual(result["jsonSchema"]["status"], "PASS")
                self.assertEqual(result["profile"]["artifactClass"], expected[0])
                self.assertEqual(result["profile"]["primaryOutput"]["format"], expected[1])

    def test_cross_format_profile_rejects_wrong_primary_format(self) -> None:
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / "bad.json"
            value = json.loads((ROOT / "artifact-delivery/examples/document-r1.json").read_text())
            value.pop("$schema", None)
            value["primaryOutput"]["format"] = "xlsx"
            path.write_text(json.dumps(value))
            result = MODULE.validate_profile(path)
            self.assertEqual(result["status"], "FAIL")
            self.assertTrue(any("document primaryOutput.format" in item for item in result["minimalContractErrors"]))

    def test_delivery_request_binds_exact_profile_and_source_bytes(self) -> None:
        request = ROOT / "artifact-delivery/examples/presentation-native-smoke-request-r1.json"
        result = MODULE.validate_delivery_request(request)
        if result["requestValidation"]["jsonSchema"]["validator"] == "unavailable":
            self.assertEqual(result["status"], "FAIL")
            self.assertIn("request envelope did not PASS schema validation", result["failures"])
        else:
            self.assertEqual(result["status"], "PASS", result)
        self.assertEqual(result["resolved"]["profile"]["digest"]["sha256"], result["request"]["profile"]["sha256"])
        self.assertEqual(result["resolved"]["source"]["digest"]["sha256"], result["request"]["source"]["sha256"])

    def test_delivery_request_detects_source_digest_drift(self) -> None:
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            profile_source = ROOT / "artifact-delivery/examples/pdu-sdu-presentation-r1.json"
            source_source = ROOT / "artifact-delivery/examples/presentation-native-smoke-source-r1.json"
            profile = root / "profile.json"
            source = root / "source.json"
            shutil.copyfile(profile_source, profile)
            shutil.copyfile(source_source, source)
            request = {
                "schemaVersion": 1,
                "kind": "artifact-delivery-request",
                "requestId": "artifact-request:digest-drift",
                "profile": {"id": "pdu-sdu-presentation-r1", "path": "profile.json", "sha256": MODULE.sha256_file(profile)},
                "source": {"kind": "presentation-source-v1", "path": "source.json", "sha256": MODULE.sha256_file(source)},
                "outputDirectory": "out",
                "builder": {
                    "id": "https://ordivon.local/builders/artifact-delivery/python-pptx-v1",
                    "buildType": "https://ordivon.local/build-types/artifact-delivery/presentation-source-v1"
                }
            }
            request_path = root / "request.json"
            request_path.write_text(json.dumps(request))
            source.write_text(source.read_text() + " ")
            result = MODULE.validate_delivery_request(request_path)
            self.assertEqual(result["status"], "FAIL")
            self.assertIn("source digest mismatch", result["failures"])

    def test_compile_delivery_plan_is_derived_and_profile_dispatched(self) -> None:
        request = ROOT / "artifact-delivery/examples/presentation-native-smoke-request-r1.json"
        plan = MODULE.compile_delivery_plan(request)
        self.assertEqual(plan["kind"], "artifact-delivery-derived-plan")
        self.assertEqual(plan["buildAdapter"], "python-pptx-presentation-source-v1")
        self.assertEqual(plan["stages"], ["build", "verify", "package", "release"])
        self.assertIn("target", plan["requiredGates"])
        self.assertIn("visual", plan["requiredGates"])
        if importlib.util.find_spec("jsonschema") is None:
            self.assertEqual(plan["status"], "FAIL")
        else:
            self.assertEqual(plan["status"], "PASS", plan)

    def test_semantic_svg_source_is_digest_bound_and_routes_to_ppt_master(self) -> None:
        if importlib.util.find_spec("jsonschema") is None:
            self.skipTest("jsonschema is unavailable")
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            profile = root / "profile.json"
            shutil.copyfile(ROOT / "artifact-delivery/examples/pdu-sdu-presentation-r1.json", profile)
            page = root / "page.svg"
            page.write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 720"/>')
            source = {
                "schemaVersion": 1,
                "kind": "presentation-semantic-svg-source",
                "presentationId": "presentation:semantic-svg-routing-smoke-r1",
                "profileId": "pdu-sdu-presentation-r1",
                "locale": "en-US",
                "transition": "none",
                "nativeChartsAndTables": False,
                "pages": [{"id": "page-01", "path": page.name, "sha256": MODULE.sha256_file(page)}],
            }
            source_path = root / "source.json"
            source_path.write_text(json.dumps(source))
            request = {
                "schemaVersion": 1,
                "kind": "artifact-delivery-request",
                "requestId": "artifact-request:semantic-svg-routing-smoke-r1",
                "profile": {"id": "pdu-sdu-presentation-r1", "path": profile.name, "sha256": MODULE.sha256_file(profile)},
                "source": {"kind": "presentation-semantic-svg-source-v1", "path": source_path.name, "sha256": MODULE.sha256_file(source_path)},
                "outputDirectory": "out",
                "builder": {
                    "id": "https://ordivon.local/builders/artifact-delivery/ppt-master-v1",
                    "buildType": "https://ordivon.local/build-types/artifact-delivery/presentation-semantic-svg-source-v1",
                },
            }
            request_path = root / "request.json"
            request_path.write_text(json.dumps(request))
            validation = MODULE.validate_delivery_request(request_path)
            self.assertEqual(validation["status"], "PASS", validation)
            self.assertEqual(len(validation["resolved"]["materials"]), 1)
            self.assertEqual(validation["resolved"]["materials"][0]["digest"]["sha256"], MODULE.sha256_file(page))
            plan = MODULE.compile_delivery_plan(request_path)
            self.assertEqual(plan["status"], "PASS", plan)
            self.assertEqual(plan["buildAdapter"], "ppt-master-semantic-svg-v1")
            self.assertEqual(plan["builder"]["id"], "https://ordivon.local/builders/artifact-delivery/ppt-master-v1")
            page.write_text(page.read_text() + "\n<!-- drift -->")
            drift = MODULE.validate_delivery_request(request_path)
            self.assertEqual(drift["status"], "FAIL")
            self.assertTrue(any("semantic SVG page digest mismatch" in item for item in drift["failures"]), drift)

    def test_semantic_svg_source_rejects_parent_traversal_material_target(self) -> None:
        if importlib.util.find_spec("jsonschema") is None:
            self.skipTest("jsonschema is unavailable")
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            page = root / "page.svg"
            page.write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 720"/>')
            asset = root / "asset.png"
            asset.write_bytes(b"not-an-image-but-digest-bound")
            source = {
                "schemaVersion": 1,
                "kind": "presentation-semantic-svg-source",
                "presentationId": "presentation:semantic-svg-path-smoke-r1",
                "profileId": "pdu-sdu-presentation-r1",
                "locale": "en-US",
                "transition": "none",
                "nativeChartsAndTables": False,
                "pages": [{"id": "page-01", "path": page.name, "sha256": MODULE.sha256_file(page)}],
                "materials": [{"path": asset.name, "sha256": MODULE.sha256_file(asset), "projectRelativePath": "../escape.png"}],
            }
            source_path = root / "source.json"
            source_path.write_text(json.dumps(source))
            result = MODULE.validate_json_document(
                source_path,
                ROOT / "artifact-delivery/presentation-semantic-svg-source-v1.schema.json",
                "presentation-semantic-svg-source",
            )
            self.assertEqual(result["status"], "FAIL", result)

    def test_ppt_master_provider_commit_drift_fails_closed(self) -> None:
        from unittest import mock
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            for relative in (
                ".venv-exp/bin/python",
                "skills/ppt-master/scripts/svg_quality_checker.py",
                "skills/ppt-master/scripts/svg_to_pptx.py",
            ):
                path = root / relative
                path.parent.mkdir(parents=True, exist_ok=True)
                path.write_text("stub")
            completed = subprocess.CompletedProcess(
                args=["git"],
                returncode=0,
                stdout="0" * 40 + "\n",
                stderr="",
            )
            with mock.patch.dict(os.environ, {"ARTIFACT_PPT_MASTER_ROOT": str(root)}, clear=False), mock.patch.object(MODULE.subprocess, "run", return_value=completed):
                provider, failures = MODULE._ppt_master_provider_facts()
            self.assertEqual(provider["observedCommit"], "0" * 40)
            self.assertTrue(any("provider commit mismatch" in item for item in failures), failures)

    def test_native_presentation_source_build_rejects_undeclared_font(self) -> None:
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            source = json.loads((ROOT / "artifact-delivery/examples/presentation-native-smoke-source-r1.json").read_text())
            source.pop("$schema", None)
            source["slides"][0]["elements"][0]["fontFamily"] = "Aptos"
            source_path = root / "source.json"
            source_path.write_text(json.dumps(source))
            output = root / "out.pptx"
            result = MODULE.build_presentation_source(
                source_path,
                ROOT / "artifact-delivery/examples/pdu-sdu-presentation-r1.json",
                output,
            )
            self.assertEqual(result["status"], "FAIL")
            self.assertTrue(any("undeclared font family" in item for item in result["failures"]))
            self.assertFalse(output.exists())

    def _template_bound_source(self, root: Path, *, add_sample_slide: bool = False) -> tuple[Path, Path]:
        if importlib.util.find_spec("pptx") is None:
            self.skipTest("python-pptx is unavailable")
        from pptx import Presentation
        from pptx.util import Inches
        template = root / "authoring-template.pptx"
        prs = Presentation()
        prs.slide_width = Inches(13.333333)
        prs.slide_height = Inches(7.5)
        if add_sample_slide:
            prs.slides.add_slide(prs.slide_layouts[6])
        prs.save(template)
        source = {
            "schemaVersion": 1,
            "kind": "presentation-source",
            "presentationId": "presentation:template-binding-smoke-r1",
            "profileId": "pdu-sdu-presentation-r1",
            "sourceMode": "native-composition",
            "locale": "zh-CN",
            "aspectRatio": "16:9",
            "slideSizeInches": {"width": 13.333333, "height": 7.5},
            "template": {"format": "pptx", "path": template.name, "sha256": MODULE.sha256_file(template)},
            "slides": [{
                "id": "slide-01",
                "title": "Template-bound title",
                "layoutId": "Title Slide",
                "elements": [
                    {"kind": "text", "id": "title", "text": "Template-bound title", "placeholderIdx": 0},
                    {"kind": "text", "id": "subtitle", "text": "Inherited layout styling", "placeholderIdx": 1}
                ]
            }]
        }
        source_path = root / "source.json"
        source_path.write_text(json.dumps(source))
        return template, source_path

    def test_presentation_template_master_layout_placeholder_binding(self) -> None:
        if importlib.util.find_spec("pptx") is None or importlib.util.find_spec("jsonschema") is None:
            self.skipTest("python-pptx/jsonschema are not available")
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            template, source = self._template_bound_source(root)
            output = root / "bound.pptx"
            result = MODULE.build_presentation_source(source, ROOT / "artifact-delivery/examples/pdu-sdu-presentation-r1.json", output)
            self.assertEqual(result["status"], "PASS", result)
            self.assertEqual(result["template"]["artifact"]["digest"]["sha256"], MODULE.sha256_file(template))
            self.assertGreaterEqual(len(result["template"]["themeParts"]), 1)
            self.assertGreaterEqual(len(result["template"]["masterParts"]), 1)
            self.assertGreaterEqual(len(result["template"]["layoutParts"]), 1)
            self.assertEqual(result["layoutBindings"][0]["layout"]["name"], "Title Slide")
            self.assertTrue(result["layoutBindings"][0]["layout"]["masterPart"].startswith("/ppt/slideMasters/"))
            output_template_fact = MODULE._presentation_template_package_fact(output)
            self.assertEqual(output_template_fact["themeParts"], result["template"]["themeParts"])
            self.assertEqual(output_template_fact["masterParts"], result["template"]["masterParts"])
            self.assertEqual(output_template_fact["layoutParts"], result["template"]["layoutParts"])
            observed = Presentation(output)
            self.assertEqual(len(observed.slides), 1)
            placeholders = {int(p.placeholder_format.idx): p for p in observed.slides[0].placeholders}
            self.assertEqual(placeholders[0].text, "Template-bound title")
            self.assertEqual(placeholders[1].text, "Inherited layout styling")
            validator = openxml_validator_path()
            dotnet = openxml_dotnet_path()
            if validator.is_file() and dotnet.is_file():
                structural = MODULE.verify_openxml_artifact(output)
                self.assertEqual(structural["status"], "PASS", structural)

    def test_presentation_template_is_promoted_to_digest_bound_build_material(self) -> None:
        if importlib.util.find_spec("pptx") is None or importlib.util.find_spec("jsonschema") is None:
            self.skipTest("python-pptx/jsonschema are not available")
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            template, source = self._template_bound_source(root)
            profile_source = ROOT / "artifact-delivery/examples/pdu-sdu-presentation-r1.json"
            profile = root / "profile.json"
            shutil.copyfile(profile_source, profile)
            request = {
                "schemaVersion": 1,
                "kind": "artifact-delivery-request",
                "requestId": "artifact-request:template-material-smoke",
                "profile": {"id": "pdu-sdu-presentation-r1", "path": profile.name, "sha256": MODULE.sha256_file(profile)},
                "source": {"kind": "presentation-source-v1", "path": source.name, "sha256": MODULE.sha256_file(source)},
                "outputDirectory": "out",
                "builder": {
                    "id": "https://ordivon.local/builders/artifact-delivery/python-pptx-v1",
                    "buildType": "https://ordivon.local/build-types/artifact-delivery/presentation-source-v1"
                }
            }
            request_path = root / "request.json"
            request_path.write_text(json.dumps(request))
            validation = MODULE.validate_delivery_request(request_path)
            self.assertEqual(validation["status"], "PASS", validation)
            materials = validation["resolved"]["materials"]
            self.assertEqual(len(materials), 1)
            self.assertEqual(materials[0]["path"], str(template.resolve()))
            self.assertEqual(materials[0]["digest"]["sha256"], MODULE.sha256_file(template))
            plan = MODULE.compile_delivery_plan(request_path)
            self.assertEqual(plan["status"], "PASS", plan)
            self.assertEqual(plan["resolvedInputs"]["materials"][0]["digest"]["sha256"], MODULE.sha256_file(template))

    def test_presentation_template_digest_drift_fails_closed(self) -> None:
        if importlib.util.find_spec("pptx") is None or importlib.util.find_spec("jsonschema") is None:
            self.skipTest("python-pptx/jsonschema are not available")
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            template, source = self._template_bound_source(root)
            value = json.loads(source.read_text())
            value["template"]["sha256"] = "0" * 64
            source.write_text(json.dumps(value))
            result = MODULE.build_presentation_source(source, ROOT / "artifact-delivery/examples/pdu-sdu-presentation-r1.json", root / "out.pptx")
            self.assertEqual(result["status"], "FAIL")
            self.assertTrue(any("template digest mismatch" in item for item in result["failures"]))

    def test_presentation_template_rejects_missing_placeholder(self) -> None:
        if importlib.util.find_spec("pptx") is None or importlib.util.find_spec("jsonschema") is None:
            self.skipTest("python-pptx/jsonschema are not available")
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            _, source = self._template_bound_source(root)
            value = json.loads(source.read_text())
            value["slides"][0]["elements"][0]["placeholderIdx"] = 999
            source.write_text(json.dumps(value))
            result = MODULE.build_presentation_source(source, ROOT / "artifact-delivery/examples/pdu-sdu-presentation-r1.json", root / "out.pptx")
            self.assertEqual(result["status"], "FAIL")
            self.assertTrue(any("placeholder idx 999 is absent" in item for item in result["failures"]))

    def test_presentation_template_rejects_sample_slide_contamination(self) -> None:
        if importlib.util.find_spec("pptx") is None or importlib.util.find_spec("jsonschema") is None:
            self.skipTest("python-pptx/jsonschema are not available")
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            _, source = self._template_bound_source(root, add_sample_slide=True)
            result = MODULE.build_presentation_source(source, ROOT / "artifact-delivery/examples/pdu-sdu-presentation-r1.json", root / "out.pptx")
            self.assertEqual(result["status"], "FAIL")
            self.assertTrue(any("must contain zero slides" in item for item in result["failures"]))

    def _hybrid_media_source(self, root: Path, *, decorative: bool = True, include_alt: bool = False) -> tuple[Path, Path]:
        if importlib.util.find_spec("PIL") is None:
            self.skipTest("Pillow is unavailable")
        from PIL import Image, ImageDraw
        image_path = root / "visual-authority.png"
        image = Image.new("RGB", (640, 360), "white")
        draw = ImageDraw.Draw(image)
        draw.rectangle((20, 20, 620, 340), outline="black", width=4)
        draw.text((48, 48), "VISUAL AUTHORITY", fill="black")
        image.save(image_path)
        element = {
            "kind": "image",
            "id": "visual-reference",
            "path": image_path.name,
            "sha256": MODULE.sha256_file(image_path),
            "box": {"x": 0, "y": 0, "w": 13.333333, "h": 7.5},
            "decorative": decorative
        }
        if include_alt:
            element["altText"] = "Approved high-fidelity visual reference"
        source = {
            "schemaVersion": 1,
            "kind": "presentation-source",
            "presentationId": "presentation:hybrid-media-smoke-r1",
            "profileId": "pdu-sdu-presentation-r1",
            "sourceMode": "native-composition",
            "locale": "zh-CN",
            "aspectRatio": "16:9",
            "slideSizeInches": {"width": 13.333333, "height": 7.5},
            "slides": [{
                "id": "slide-01",
                "title": "Hybrid media smoke",
                "layoutId": "Blank",
                "elements": [
                    element,
                    {
                        "kind": "text",
                        "id": "native-title",
                        "text": "Native editable overlay",
                        "box": {"x": 0.8, "y": 0.5, "w": 5.0, "h": 0.6},
                        "fontFamily": "Arial",
                        "fontSizePt": 24,
                        "bold": True,
                        "colorHex": "111111"
                    }
                ]
            }]
        }
        source_path = root / "source.json"
        source_path.write_text(json.dumps(source))
        return image_path, source_path

    def test_hybrid_presentation_binds_exact_raster_and_native_overlay(self) -> None:
        if importlib.util.find_spec("pptx") is None or importlib.util.find_spec("jsonschema") is None:
            self.skipTest("python-pptx/jsonschema are not available")
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            image_path, source = self._hybrid_media_source(root, decorative=False, include_alt=True)
            output = root / "hybrid.pptx"
            result = MODULE.build_presentation_source(source, ROOT / "artifact-delivery/examples/pdu-sdu-presentation-r1.json", output)
            self.assertEqual(result["status"], "PASS", result)
            self.assertEqual(len(result["mediaBindings"]), 1)
            self.assertEqual(result["mediaBindings"][0]["artifact"]["digest"]["sha256"], MODULE.sha256_file(image_path))
            observed = Presentation(output)
            self.assertEqual(len(observed.slides), 1)
            self.assertEqual(observed.slides[0].shapes[0].shape_type, MSO_SHAPE_TYPE.PICTURE)
            self.assertTrue(observed.slides[0].shapes[1].has_text_frame)
            self.assertEqual(observed.slides[0].shapes[1].text, "Native editable overlay")
            c_nv_pr = observed.slides[0].shapes[0]._element.xpath('.//p:cNvPr')[0]
            self.assertEqual(c_nv_pr.get('descr'), "Approved high-fidelity visual reference")
            with zipfile.ZipFile(output) as package:
                media = [name for name in package.namelist() if name.startswith("ppt/media/")]
                self.assertEqual(len(media), 1)
                self.assertEqual(MODULE.hashlib.sha256(package.read(media[0])).hexdigest(), MODULE.sha256_file(image_path))
            validator = openxml_validator_path()
            dotnet = openxml_dotnet_path()
            if validator.is_file() and dotnet.is_file():
                structural = MODULE.verify_openxml_artifact(output)
                self.assertEqual(structural["status"], "PASS", structural)

    def test_ultrawide_34x10_hybrid_profile_builds_native_pptx(self) -> None:
        if importlib.util.find_spec("pptx") is None or importlib.util.find_spec("jsonschema") is None or importlib.util.find_spec("PIL") is None:
            self.skipTest("python-pptx/jsonschema/Pillow are not available")
        from PIL import Image
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            image_path = root / "ultrawide-reference.png"
            Image.new("RGB", (680, 200), "white").save(image_path)
            source = {
                "schemaVersion": 1, "kind": "presentation-source",
                "presentationId": "presentation:ultrawide-hybrid-smoke-r1",
                "profileId": "presentation-ultrawide-34x10-r1",
                "sourceMode": "native-composition", "locale": "zh-CN", "aspectRatio": "34:10",
                "slideSizeInches": {"width": 34.0, "height": 10.0},
                "slides": [{"id": "slide-01", "title": "Ultrawide hybrid", "layoutId": "Blank", "elements": [
                    {"kind": "image", "id": "reference", "path": image_path.name, "sha256": MODULE.sha256_file(image_path), "box": {"x": 0, "y": 0, "w": 34.0, "h": 10.0}, "decorative": True},
                    {"kind": "text", "id": "native-title", "text": "34:10 native overlay", "box": {"x": 1.2, "y": 0.7, "w": 10.0, "h": 0.8}, "fontFamily": "Arial", "fontSizePt": 28, "bold": True}
                ]}]
            }
            source_path = root / "source.json"
            source_path.write_text(json.dumps(source))
            profile = ROOT / "artifact-delivery/examples/presentation-ultrawide-34x10-r1.json"
            self.assertEqual(MODULE.validate_profile(profile)["status"], "PASS")
            output = root / "ultrawide.pptx"
            result = MODULE.build_presentation_source(source_path, profile, output)
            self.assertEqual(result["status"], "PASS", result)
            self.assertEqual(result["semantic"]["status"], "PASS", result)
            observed = Presentation(output)
            self.assertAlmostEqual(float(observed.slide_width) / float(observed.slide_height), 3.4, places=5)
            validator = openxml_validator_path()
            dotnet = openxml_dotnet_path()
            if validator.is_file() and dotnet.is_file():
                structural = MODULE.verify_openxml_artifact(output)
                self.assertEqual(structural["status"], "PASS", structural)

    def test_hybrid_image_is_promoted_to_digest_bound_build_material(self) -> None:
        if importlib.util.find_spec("pptx") is None or importlib.util.find_spec("jsonschema") is None:
            self.skipTest("python-pptx/jsonschema are not available")
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            image_path, source = self._hybrid_media_source(root)
            profile = root / "profile.json"
            shutil.copyfile(ROOT / "artifact-delivery/examples/pdu-sdu-presentation-r1.json", profile)
            request = {
                "schemaVersion": 1, "kind": "artifact-delivery-request", "requestId": "artifact-request:hybrid-media-material-smoke",
                "profile": {"id": "pdu-sdu-presentation-r1", "path": profile.name, "sha256": MODULE.sha256_file(profile)},
                "source": {"kind": "presentation-source-v1", "path": source.name, "sha256": MODULE.sha256_file(source)},
                "outputDirectory": "out",
                "builder": {"id": "https://ordivon.local/builders/artifact-delivery/python-pptx-v1", "buildType": "https://ordivon.local/build-types/artifact-delivery/presentation-source-v1"}
            }
            request_path = root / "request.json"
            request_path.write_text(json.dumps(request))
            validation = MODULE.validate_delivery_request(request_path)
            self.assertEqual(validation["status"], "PASS", validation)
            self.assertEqual(len(validation["resolved"]["materials"]), 1)
            self.assertEqual(validation["resolved"]["materials"][0]["path"], str(image_path.resolve()))
            self.assertEqual(validation["resolved"]["materials"][0]["digest"]["sha256"], MODULE.sha256_file(image_path))

    def test_hybrid_image_digest_drift_fails_closed(self) -> None:
        if importlib.util.find_spec("pptx") is None or importlib.util.find_spec("jsonschema") is None:
            self.skipTest("python-pptx/jsonschema are not available")
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            _, source = self._hybrid_media_source(root)
            value = json.loads(source.read_text())
            value["slides"][0]["elements"][0]["sha256"] = "0" * 64
            source.write_text(json.dumps(value))
            result = MODULE.build_presentation_source(source, ROOT / "artifact-delivery/examples/pdu-sdu-presentation-r1.json", root / "out.pptx")
            self.assertEqual(result["status"], "FAIL")
            self.assertTrue(any("image digest mismatch" in item for item in result["failures"]))

    def test_non_decorative_hybrid_image_requires_alt_text(self) -> None:
        if importlib.util.find_spec("jsonschema") is None:
            self.skipTest("jsonschema is unavailable")
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            _, source = self._hybrid_media_source(root, decorative=False, include_alt=False)
            result = MODULE.validate_json_document(source, ROOT / "artifact-delivery/presentation-source-v1.schema.json", "presentation-source")
            self.assertEqual(result["status"], "FAIL")
            self.assertEqual(result["jsonSchema"]["status"], "FAIL")

    def test_transparent_native_text_overlay_uses_standard_drawingml_alpha(self) -> None:
        if importlib.util.find_spec("pptx") is None or importlib.util.find_spec("jsonschema") is None:
            self.skipTest("python-pptx/jsonschema are not available")
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            source = {
                "schemaVersion": 1, "kind": "presentation-source",
                "presentationId": "presentation:transparent-overlay-smoke-r1",
                "profileId": "presentation-ultrawide-34x10-r1",
                "sourceMode": "native-composition", "locale": "zh-CN", "aspectRatio": "34:10",
                "slideSizeInches": {"width": 34.0, "height": 10.0},
                "slides": [{"id": "slide-01", "title": "Transparent overlay", "layoutId": "Blank", "elements": [{
                    "kind": "text", "id": "semantic-title", "text": "面向敏捷交付的\n研发组织变革",
                    "box": {"x": 1.0, "y": 1.0, "w": 10.0, "h": 1.0},
                    "fontFamily": "Microsoft YaHei", "fontSizePt": 32, "bold": True,
                    "colorHex": "0B2F74", "opacity": 0.0
                }]}]
            }
            source_path = root / "source.json"
            source_path.write_text(json.dumps(source))
            output = root / "transparent.pptx"
            result = MODULE.build_presentation_source(source_path, ROOT / "artifact-delivery/examples/presentation-ultrawide-34x10-r1.json", output)
            self.assertEqual(result["status"], "PASS", result)
            with zipfile.ZipFile(output) as package:
                xml = package.read("ppt/slides/slide1.xml").decode("utf-8")
                self.assertEqual(xml.count('a:alpha val="0"'), 2)
                self.assertIn("面向敏捷交付的", xml)
                self.assertIn("研发组织变革", xml)
            structural = MODULE.verify_openxml_artifact(output)
            if structural.get("status") != "NOT_RUN":
                self.assertEqual(structural["status"], "PASS", structural)

    def test_text_opacity_requires_explicit_color(self) -> None:
        if importlib.util.find_spec("jsonschema") is None:
            self.skipTest("jsonschema is unavailable")
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            source = {
                "schemaVersion": 1, "kind": "presentation-source",
                "presentationId": "presentation:transparent-overlay-negative-r1",
                "profileId": "presentation-ultrawide-34x10-r1",
                "sourceMode": "native-composition", "locale": "zh-CN", "aspectRatio": "34:10",
                "slideSizeInches": {"width": 34.0, "height": 10.0},
                "slides": [{"id": "slide-01", "title": "Negative", "layoutId": "Blank", "elements": [{
                    "kind": "text", "id": "semantic-title", "text": "x",
                    "box": {"x": 1.0, "y": 1.0, "w": 2.0, "h": 1.0},
                    "fontFamily": "Arial", "fontSizePt": 24, "opacity": 0.0
                }]}]
            }
            source_path = root / "source.json"
            source_path.write_text(json.dumps(source))
            result = MODULE.build_presentation_source(source_path, ROOT / "artifact-delivery/examples/presentation-ultrawide-34x10-r1.json", root / "out.pptx")
            self.assertEqual(result["status"], "FAIL")
            self.assertTrue(any("opacity requires explicit colorHex" in item for item in result["failures"]))

    def test_pdu_sdu_34x10_slide01_reference_map_validates_but_is_not_native_build(self) -> None:
        source = ROOT / "artifact-delivery/golden/pdu-sdu-34x10-r1/pdu-sdu-34x10-slide01-reference-map-r1.json"
        validation = MODULE.validate_json_document(source, ROOT / "artifact-delivery/presentation-source-v1.schema.json", "presentation-source")
        if importlib.util.find_spec("jsonschema") is None:
            self.assertEqual(validation["status"], "FAIL")
        else:
            self.assertEqual(validation["status"], "PASS", validation)
        built = MODULE.build_presentation_source(source, ROOT / "artifact-delivery/examples/presentation-ultrawide-34x10-r1.json", Path(tempfile.gettempdir()) / "should-not-build-reference-map.pptx")
        self.assertEqual(built["status"], "FAIL")
        self.assertTrue(any("sourceMode=native-composition only" in item for item in built["failures"]))

    def test_pdu_sdu_34x10_deck_reference_map_binds_all_eight_visual_authorities(self) -> None:
        source = ROOT / "artifact-delivery/golden/pdu-sdu-34x10-r1/pdu-sdu-34x10-deck-reference-map-r1.json"
        validation = MODULE.validate_json_document(source, ROOT / "artifact-delivery/presentation-source-v1.schema.json", "presentation-source")
        if importlib.util.find_spec("jsonschema") is None:
            self.assertEqual(validation["status"], "FAIL")
            return
        self.assertEqual(validation["status"], "PASS", validation)
        doc = validation["document"]
        self.assertEqual(doc["sourceMode"], "reference-map")
        self.assertEqual(doc["profileId"], "presentation-ultrawide-34x10-r1")
        self.assertEqual(doc["aspectRatio"], "34:10")
        self.assertEqual(len(doc["slides"]), 8)
        expected = [
            "97237ff50567e495b0ec56543dd644143fed4a5241dd038814c775c732ceb998",
            "fc9e38e60aa35f36050cef2d0493d716ebc20977b13a126bdd1560dc6643b168",
            "bd511f6736bf5c8eee439c412e4e32fc7f0ca57953da6baee4a9e716c86cdec9",
            "116de845c183109a27c68fb2ada0c5f1038ea55ea5b4f682bf37dee3017fc875",
            "23c7b1a818feb6a8e7de9f009f57542acffea306e22fd7a83eb6c10063340d65",
            "2eb714c894a4115fcd3e55a0729d28ddd93e3e92ec2621c0274a0b6131bbfeb5",
            "67287373393b0fb170b0d0f3957661115a01a4e9cd43655c65d344729708379a",
            "f1727f930369788b5c6a314c4fe2d3baf9f4c27e2eb361ae8ef2ff2964ef1628",
        ]
        self.assertEqual([slide["legacySource"]["sha256"] for slide in doc["slides"]], expected)
        built = MODULE.build_presentation_source(source, ROOT / "artifact-delivery/examples/presentation-ultrawide-34x10-r1.json", Path(tempfile.gettempdir()) / "should-not-build-deck-reference-map.pptx")
        self.assertEqual(built["status"], "FAIL")
        self.assertTrue(any("sourceMode=native-composition only" in item for item in built["failures"]))

    def test_pdu_sdu_34x10_deck_migration_evidence_has_no_business_validation_claim(self) -> None:
        evidence = json.loads((ROOT / "artifact-delivery/golden/pdu-sdu-34x10-r1/pdu-sdu-34x10-deck-migration-evidence-r1.json").read_text())
        self.assertEqual(len(evidence["slides"]), 8)
        self.assertEqual(evidence["currentNativeMigration"]["standing"], "DECK_SEMANTIC_OVERLAY_NATIVE_BUILD_PASS_FINAL_HYBRID_NOT_RUN")
        self.assertIn("does not validate or alter business/organizational claims", evidence["boundary"])
        self.assertIn("BLOCKED_ON_EXACT_RASTER_TRANSPORT", evidence["migrationStanding"])

    def test_pdu_sdu_34x10_r2_acceptance_manifest_binds_exact_accepted_tuple(self) -> None:
        path = ROOT / "artifact-delivery/golden/pdu-sdu-34x10-r1/pdu-sdu-34x10-acceptance-r2.json"
        value = json.loads(path.read_text())
        self.assertEqual(value["kind"], "artifact-delivery-golden-acceptance")
        self.assertEqual(value["standing"], "ACCEPTED_IN_PROFILE")
        self.assertEqual(value["profile"]["id"], "presentation-ultrawide-34x10-r1")
        self.assertEqual(value["profile"]["sha256"], "56cf5879e52984fb398242416bb46f1670e9aecd7e5d88e18677b4263126eea9")
        self.assertEqual(value["sourceAuthority"]["sha256"], "b50ab3025c4b285c728d029416492b4c7a03de1fc9e0bfe3f05293bfaf5f468c")
        self.assertEqual(value["sourceAuthority"]["admission"], "workspace.execBound")
        self.assertEqual(value["sourceCurrentness"]["buildSourceRevision"], "cdc0087b208a6009717383f7b58747f2f310895b")
        self.assertEqual(value["sourceCurrentness"]["presentationCriticalOverlap"], [])
        self.assertEqual(value["acceptedTuple"]["primary"]["sha256"], "fa48aa4a513f7ab04e05cc374b80592903acae8c39d727cf3c9737d68f2f1d03")
        self.assertEqual(value["acceptedTuple"]["primary"]["sizeBytes"], 2745582)
        self.assertEqual(value["acceptedTuple"]["primary"]["slideCount"], 8)
        self.assertEqual(value["acceptedTuple"]["companion"]["sha256"], "bea385bfc2d78b82498315c251203c4ee54f0112fe24430a8142cd1aebb850fe")
        self.assertEqual(value["acceptedTuple"]["companion"]["sizeBytes"], 1149721)
        self.assertEqual(value["structuralEvidence"]["openXmlValidator"]["validationErrorCount"], 0)
        self.assertEqual(value["targetEvidence"]["renderer"], "Microsoft PowerPoint Desktop")
        self.assertTrue(value["targetEvidence"]["sourceDigestStableAcrossRender"])
        self.assertEqual(len(value["targetEvidence"]["renderPngSha256"]), 8)
        self.assertEqual(value["visualEvidence"]["blockingDefects"], [])
        self.assertEqual(value["deliveryEvidence"]["windows-workstation"]["status"], "PASS")
        self.assertEqual(value["deliveryEvidence"]["google-drive"]["status"], "PASS")
        self.assertEqual(value["deliveryEvidence"]["google-drive"]["primaryFileId"], "external-gdrive:file:1AZ9qdZpix0yxGioi17dTFiMq4USdoFoY")
        self.assertEqual(value["deliveryEvidence"]["google-drive"]["companionFileId"], "external-gdrive:file:1t2c4akOyk1YGlVv6ZcpAgMG1h44AwuVr")
        self.assertEqual(value["finalGate"]["status"], "PASS")
        self.assertEqual(value["finalGate"]["requiredGateFailures"], [])
        self.assertEqual(value["finalGate"]["evidenceSha256"], "c3e9666b287f7dc29d1eb6371f1aaa088c0eb56b9708e3d29658f9ca29feeecd")
        nonclaims = set(value["nonClaims"])
        self.assertIn("BUSINESS_CONTENT_CORRECTNESS_NOT_ESTABLISHED", nonclaims)
        self.assertIn("ACCESSIBILITY_AND_HUMAN_USABILITY_NOT_ACCEPTED", nonclaims)
        self.assertIn("CROSS_WORKSPACE_BYTE_DETERMINISM_NOT_PROVEN", nonclaims)
        self.assertIn("FORMAL_THIRD_PARTY_INDEPENDENT_IVV_NOT_CLAIMED", nonclaims)
        self.assertEqual(value["openFaults"][0]["standing"], "OPEN_SEPARATE_FROM_ACCEPTED_EXACT_TUPLE")

    def test_pdu_sdu_34x10_runtime_input_set_binds_exact_eight_visuals(self) -> None:
        value = json.loads((ROOT / "artifact-delivery/golden/pdu-sdu-34x10-r1/pdu-sdu-34x10-runtime-input-set-r1.json").read_text())
        self.assertEqual(value["kind"], "artifact-delivery-runtime-input-set")
        self.assertEqual(value["runtimeAuthority"], "artifact-golden-r1")
        self.assertEqual(value["admission"], "workspace.execBound")
        self.assertEqual(value["standing"], "AUTHORITY_CONFIGURED_OBJECTS_NOT_YET_MATERIALIZED")
        self.assertEqual(len(value["bindings"]), 8)
        expected = [
            "97237ff50567e495b0ec56543dd644143fed4a5241dd038814c775c732ceb998",
            "fc9e38e60aa35f36050cef2d0493d716ebc20977b13a126bdd1560dc6643b168",
            "bd511f6736bf5c8eee439c412e4e32fc7f0ca57953da6baee4a9e716c86cdec9",
            "116de845c183109a27c68fb2ada0c5f1038ea55ea5b4f682bf37dee3017fc875",
            "23c7b1a818feb6a8e7de9f009f57542acffea306e22fd7a83eb6c10063340d65",
            "2eb714c894a4115fcd3e55a0729d28ddd93e3e92ec2621c0274a0b6131bbfeb5",
            "67287373393b0fb170b0d0f3957661115a01a4e9cd43655c65d344729708379a",
            "f1727f930369788b5c6a314c4fe2d3baf9f4c27e2eb361ae8ef2ff2964ef1628",
        ]
        self.assertEqual([b["expectedDigest"] for b in value["bindings"]], [f"sha256:{x}" for x in expected])
        self.assertEqual([b["presentationRelativePath"] for b in value["bindings"]], [f"visual/slide-{i:02d}.jpeg" for i in range(1, 9)])
        self.assertTrue(all(b["relativeObject"].startswith("pdu-sdu/34x10/") for b in value["bindings"]))

    def test_pdu_sdu_34x10_deck_semantic_overlay_builds_eight_native_slides(self) -> None:
        if importlib.util.find_spec("pptx") is None or importlib.util.find_spec("jsonschema") is None:
            self.skipTest("python-pptx/jsonschema are not available")
        source = ROOT / "artifact-delivery/golden/pdu-sdu-34x10-r1/pdu-sdu-34x10-deck-semantic-overlay-r1.json"
        profile = ROOT / "artifact-delivery/examples/presentation-ultrawide-34x10-r1.json"
        self.assertEqual(MODULE.sha256_file(source), "5ed264937da812c2fe67f4a13566d4bdc4dbac3c700a6dbdcbed4bff9f6f8743")
        with tempfile.TemporaryDirectory() as d:
            output = Path(d) / "deck-overlay.pptx"
            result = MODULE.build_presentation_source(source, profile, output)
            self.assertEqual(result["status"], "PASS", result)
            self.assertEqual(result["slideCount"], 8)
            alpha_count = 0
            text_run_count = 0
            with zipfile.ZipFile(output) as package:
                for i in range(1, 9):
                    xml = package.read(f"ppt/slides/slide{i}.xml").decode("utf-8")
                    alpha_count += xml.count('a:alpha val="0"')
                    text_run_count += xml.count("<a:r>")
            self.assertEqual(alpha_count, text_run_count)
            self.assertGreater(alpha_count, 111)
            structural = MODULE.verify_openxml_artifact(output)
            if structural.get("status") != "NOT_RUN":
                self.assertEqual(structural["status"], "PASS", structural)

    def test_zip_timestamp_normalization_changes_only_container_time_metadata(self) -> None:
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            a = root / "a.zip"
            b = root / "b.zip"
            payloads = [("alpha.txt", b"alpha" * 100), ("nested/beta.bin", bytes(range(64)))]
            for path, stamp in ((a, (2026, 9, 11, 10, 7, 2)), (b, (2026, 9, 11, 10, 7, 20))):
                with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as package:
                    for name, data in payloads:
                        info = zipfile.ZipInfo(name, stamp)
                        info.compress_type = zipfile.ZIP_DEFLATED
                        package.writestr(info, data)
            self.assertNotEqual(a.read_bytes(), b.read_bytes())
            result_a = MODULE.normalize_zip_member_timestamps(a)
            result_b = MODULE.normalize_zip_member_timestamps(b)
            self.assertEqual(result_a["status"], "PASS")
            self.assertFalse(result_a["memberPayloadBytesChanged"])
            self.assertFalse(result_a["recompressionPerformed"])
            self.assertEqual(a.read_bytes(), b.read_bytes())
            with zipfile.ZipFile(a) as package:
                self.assertEqual({tuple(info.date_time) for info in package.infolist()}, {MODULE.DETERMINISTIC_ZIP_DATETIME})
                self.assertEqual([(n, package.read(n)) for n, _ in payloads], payloads)

    def test_native_presentation_source_build_is_byte_deterministic(self) -> None:
        if importlib.util.find_spec("pptx") is None or importlib.util.find_spec("jsonschema") is None:
            self.skipTest("python-pptx/jsonschema are not available in this Python environment")
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            source = ROOT / "artifact-delivery/examples/presentation-native-smoke-source-r1.json"
            profile = ROOT / "artifact-delivery/examples/pdu-sdu-presentation-r1.json"
            first = root / "first.pptx"
            second = root / "second.pptx"
            one = MODULE.build_presentation_source(source, profile, first)
            two = MODULE.build_presentation_source(source, profile, second)
            self.assertEqual(one["status"], "PASS", one)
            self.assertEqual(two["status"], "PASS", two)
            self.assertEqual(first.read_bytes(), second.read_bytes())
            self.assertEqual(one["artifact"]["digest"]["sha256"], two["artifact"]["digest"]["sha256"])
            self.assertEqual(one["containerNormalization"]["method"], "in-place-local-and-central-dos-timestamp-normalization")
            self.assertFalse(one["containerNormalization"]["memberPayloadBytesChanged"])
            with zipfile.ZipFile(first) as package:
                self.assertEqual({tuple(info.date_time) for info in package.infolist()}, {MODULE.DETERMINISTIC_ZIP_DATETIME})

    def test_native_presentation_source_builds_ooxml_when_dependencies_available(self) -> None:
        if importlib.util.find_spec("pptx") is None or importlib.util.find_spec("jsonschema") is None:
            self.skipTest("python-pptx/jsonschema are not available in this Python environment")
        with tempfile.TemporaryDirectory() as d:
            output = Path(d) / "source-built.pptx"
            result = MODULE.build_presentation_source(
                ROOT / "artifact-delivery/examples/presentation-native-smoke-source-r1.json",
                ROOT / "artifact-delivery/examples/pdu-sdu-presentation-r1.json",
                output,
            )
            self.assertEqual(result["status"], "PASS", result)
            self.assertTrue(output.is_file())
            self.assertEqual(result["slideCount"], 2)
            self.assertEqual(result["inspection"]["status"], "PASS")
            self.assertEqual(result["semantic"]["status"], "PASS")

    def test_verification_summary_binds_subject_and_profile_without_false_slsa_claim(self) -> None:
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            subject = root / "artifact.bin"
            subject.write_bytes(b"artifact")
            profile = ROOT / "artifact-delivery/examples/pdu-sdu-presentation-r1.json"
            statement = MODULE.verification_summary_statement(
                subject,
                profile,
                "https://example.test/verifier",
                {"component": "1.0"},
                True,
            )
            statement_path = root / "statement.json"
            statement_path.write_text(json.dumps(statement))
            result = MODULE.verify_verification_summary(
                statement_path,
                subject,
                profile,
                ["https://example.test/verifier"],
            )
            self.assertEqual(result["status"], "PASS", result)
            self.assertEqual(result["verificationResult"], "PASSED")
            self.assertEqual(result["authenticity"], "NOT_VERIFIED")
            self.assertEqual(statement["predicate"]["verifiedLevels"], ["SLSA_BUILD_LEVEL_UNEVALUATED"])
            self.assertTrue(statement["predicate"]["resourceUri"].startswith("ni:///sha-256;"))
            self.assertEqual(statement["predicate"]["resourceUri"], MODULE.ni_sha256_uri(subject))
            self.assertNotIn("dependencyLevels", statement["predicate"])
            self.assertNotIn("inputAttestations", statement["predicate"])

    def test_verification_summary_detects_subject_drift(self) -> None:
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            subject = root / "artifact.bin"
            subject.write_bytes(b"artifact")
            profile = ROOT / "artifact-delivery/examples/pdu-sdu-presentation-r1.json"
            statement_path = root / "statement.json"
            statement_path.write_text(json.dumps(MODULE.verification_summary_statement(
                subject, profile, "https://example.test/verifier", {}, True
            )))
            subject.write_bytes(b"changed")
            result = MODULE.verify_verification_summary(statement_path, subject, profile)
            self.assertEqual(result["status"], "FAIL")
            self.assertTrue(any("subject does not exactly bind" in item for item in result["failures"]))

    def test_verification_summary_rejects_raw_evidence_misuse_as_input_attestation(self) -> None:
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            subject = root / "artifact.bin"
            subject.write_bytes(b"artifact")
            profile = ROOT / "artifact-delivery/examples/pdu-sdu-presentation-r1.json"
            statement = MODULE.verification_summary_statement(
                subject, profile, "https://example.test/verifier", {}, True
            )
            statement["predicate"]["inputAttestations"] = [{
                "uri": "file:///tmp/raw-validator-output.json",
                "digest": {"sha256": "0" * 64},
            }]
            statement_path = root / "statement.json"
            statement_path.write_text(json.dumps(statement))
            result = MODULE.verify_verification_summary(statement_path, subject, profile)
            self.assertEqual(result["status"], "FAIL")
            self.assertTrue(any("must not misuse inputAttestations" in item for item in result["failures"]))

    def test_vsa_gate_aggregation_fails_closed_without_authenticity(self) -> None:
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            subject = root / "artifact.pptx"
            subject.write_bytes(b"not-a-real-pptx")
            profile = ROOT / "artifact-delivery/examples/pdu-sdu-presentation-r1.json"
            vsa = root / "structural.vsa.json"
            vsa.write_text(json.dumps(MODULE.verification_summary_statement(
                subject, profile, "https://example.test/verifier", {}, True
            )))
            result = MODULE.aggregate_vsa_gates(profile, subject, {"structural": vsa})
            self.assertEqual(result["status"], "FAIL")
            self.assertTrue(any("authenticity" in item for item in result["failures"]))
            self.assertTrue(any("required gate VSA missing" in item for item in result["failures"]))

    def test_verify_stage_presentation_emits_standard_bound_vsa_receipts(self) -> None:
        if importlib.util.find_spec("pptx") is None or importlib.util.find_spec("jsonschema") is None:
            self.skipTest("python-pptx/jsonschema are not available in this Python environment")
        validator = openxml_validator_path()
        dotnet = openxml_dotnet_path()
        if not validator.is_file() or not dotnet.is_file():
            self.skipTest("Open XML validator runtime is not available")
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            pptx = root / "artifact.pptx"
            built = MODULE.build_presentation_source(
                ROOT / "artifact-delivery/examples/presentation-native-smoke-source-r1.json",
                ROOT / "artifact-delivery/examples/pdu-sdu-presentation-r1.json",
                pptx,
            )
            self.assertEqual(built["status"], "PASS", built)
            verify_dir = root / "verify"
            result = MODULE.execute_verify_stage(
                ROOT / "artifact-delivery/examples/pdu-sdu-presentation-r1.json", pptx, verify_dir
            )
            self.assertEqual(result["status"], "PASS", result)
            self.assertFalse(result["profileVerificationComplete"])
            self.assertTrue({"profileSchema", "structural", "semantic"}.issubset(result["receipts"]))
            for gate in ("profileSchema", "structural", "semantic"):
                self.assertEqual(result["receipts"][gate]["verificationResult"], "PASSED")
                self.assertEqual(result["receipts"][gate]["vsaValidation"]["status"], "PASS")

    def test_nu_html_checker_valid_fixture_when_installed(self) -> None:
        jar = ROOT / ".cache/artifact-toolchain/vnu/vnu.jar"
        if not jar.is_file() or shutil.which("java") is None:
            self.skipTest("Nu Html Checker is not installed")
        with tempfile.TemporaryDirectory() as d:
            html = Path(d) / "valid.html"
            html.write_text('<!doctype html><html lang="en"><head><meta charset="utf-8"><title>T</title></head><body><main><h1>T</h1></main></body></html>')
            result = MODULE.verify_html_conformance(html)
            self.assertEqual(result["status"], "PASS", result)
            self.assertEqual(result["messageCount"], 0)
            self.assertTrue(str(result["validator"]["version"]).startswith("26.9.7"))

    def _require_cosign_crypto(self) -> Path:
        if importlib.util.find_spec("jsonschema") is None:
            self.skipTest("jsonschema is not available")
        tool = MODULE.cosign_tool_fact()
        if tool.get("status") != "PASS":
            self.skipTest("cosign is not available")
        version = tuple(tool.get("versionTuple") or [])
        if not MODULE._version_at_least(version, MODULE.COSIGN_STANDARD_BUNDLE_MIN_VERSION):
            self.skipTest("cosign is below the standardized-bundle verification safety floor")
        return Path(tool["path"])

    def _cosign_key_and_policy(self, root: Path, verifier_id: str = MODULE.LOCAL_VSA_VERIFIER_ID) -> tuple[Path, Path, Path, dict[str, str]]:
        cosign = self._require_cosign_crypto()
        password = "artifact-e2e-test-only-password"
        env = dict(os.environ)
        env["COSIGN_PASSWORD"] = password
        key_prefix = root / "trusted-signer"
        proc = subprocess.run(
            [str(cosign), "generate-key-pair", "--output-key-prefix", str(key_prefix)],
            env=env,
            text=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            check=False,
            timeout=30,
        )
        self.assertEqual(proc.returncode, 0, proc.stderr)
        private_key = Path(str(key_prefix) + ".key")
        public_key = Path(str(key_prefix) + ".pub")
        policy = {
            "policyVersion": 1,
            "id": "test-vsa-trust-r1",
            "acceptedBundleMediaTypes": [MODULE.SIGSTORE_BUNDLE_V03],
            "signers": [{
                "id": "release-signer",
                "mode": "public-key",
                "allowedVerifierIds": [verifier_id],
                "requireTransparencyLog": False,
                "publicKey": {"path": public_key.name, "sha256": MODULE.sha256_file(public_key)},
            }],
        }
        policy_path = root / "vsa-trust-policy.json"
        policy_path.write_text(json.dumps(policy))
        signing_config = root / "private-signing-config.json"
        signing_config.write_text(json.dumps({
            "mediaType": "application/vnd.dev.sigstore.signingconfig.v0.2+json",
            "caUrls": [],
            "oidcUrls": [],
            "rekorTlogUrls": [],
            "tsaUrls": [],
        }))
        return private_key, policy_path, signing_config, env

    def _sign_vsa_no_tlog(
        self,
        subject: Path,
        statement: Path,
        bundle: Path,
        private_key: Path,
        signing_config: Path,
        env: dict[str, str],
    ) -> None:
        cosign = self._require_cosign_crypto()
        proc = subprocess.run(
            [
                str(cosign),
                "attest-blob",
                "--yes",
                "--signing-config", str(signing_config),
                "--key", str(private_key),
                "--statement", str(statement),
                "--bundle", str(bundle),
                str(subject),
            ],
            env=env,
            text=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            check=False,
            timeout=30,
        )
        self.assertEqual(proc.returncode, 0, proc.stderr)
        value = json.loads(bundle.read_text())
        self.assertEqual(value["mediaType"], MODULE.SIGSTORE_BUNDLE_V03)
        self.assertEqual(value.get("verificationMaterial", {}).get("tlogEntries", []), [])
        self.assertIn("dsseEnvelope", value)

    def test_vsa_trust_policy_binds_public_key_digest(self) -> None:
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            _, policy_path, _, _ = self._cosign_key_and_policy(root)
            result = MODULE.validate_vsa_trust_policy(policy_path)
            self.assertEqual(result["status"], "PASS", result)
            public_key = root / "trusted-signer.pub"
            public_key.write_text(public_key.read_text() + "\n")
            drift = MODULE.validate_vsa_trust_policy(policy_path)
            self.assertEqual(drift["status"], "FAIL")
            self.assertTrue(any("public key digest mismatch" in item for item in drift["failures"]))

    def test_vsa_trust_policy_binds_custom_trusted_root_digest(self) -> None:
        if importlib.util.find_spec("jsonschema") is None:
            self.skipTest("jsonschema is not available")
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            trusted_root = root / "trusted-root.json"
            trusted_root.write_text("{}")
            policy_path = root / "vsa-trust-policy.json"
            policy_path.write_text(json.dumps({
                "policyVersion": 1,
                "id": "test-keyless-root-r1",
                "acceptedBundleMediaTypes": [MODULE.SIGSTORE_BUNDLE_V03],
                "signers": [{
                    "id": "keyless-release-signer",
                    "mode": "keyless",
                    "allowedVerifierIds": [MODULE.LOCAL_VSA_VERIFIER_ID],
                    "requireTransparencyLog": True,
                    "certificateIdentity": "https://example.test/release-workflow",
                    "certificateOidcIssuer": "https://issuer.example.test",
                    "trustedRoot": {"path": trusted_root.name, "sha256": MODULE.sha256_file(trusted_root)},
                }],
            }))
            result = MODULE.validate_vsa_trust_policy(policy_path)
            self.assertEqual(result["status"], "PASS", result)
            trusted_root.write_text('{"tampered":true}')
            drift = MODULE.validate_vsa_trust_policy(policy_path)
            self.assertEqual(drift["status"], "FAIL")
            self.assertTrue(any("trustedRoot digest mismatch" in item for item in drift["failures"]))

    def test_signed_vsa_public_key_standard_bundle_verifies(self) -> None:
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            subject = root / "artifact.bin"
            subject.write_bytes(b"artifact")
            profile = ROOT / "artifact-delivery/examples/pdu-sdu-presentation-r1.json"
            statement_path = root / "structural.vsa.json"
            MODULE.write_json(statement_path, MODULE.verification_summary_statement(
                subject, profile, MODULE.LOCAL_VSA_VERIFIER_ID, {"test": "1"}, True
            ))
            private_key, policy_path, signing_config, env = self._cosign_key_and_policy(root)
            bundle = root / "structural.sigstore.json"
            self._sign_vsa_no_tlog(subject, statement_path, bundle, private_key, signing_config, env)
            result = MODULE.verify_signed_verification_summary(
                statement_path, bundle, subject, profile, policy_path, "release-signer"
            )
            self.assertEqual(result["status"], "PASS", result)
            self.assertEqual(result["authenticity"], "VERIFIED")
            self.assertEqual(result["bundleShape"]["tlogEntryCount"], 0)
            self.assertTrue(result["bundleShape"]["signedStatementMatches"])
            self.assertEqual(result["cosign"]["status"], "PASS")
            self.assertEqual(result["tool"]["provenance"]["status"], "PASS")
            self.assertEqual(result["tool"]["provenance"]["embeddedBinarySha256"], result["tool"]["sha256"])

            original_package = MODULE.COSIGN_ARCH_PACKAGE
            original_signature = MODULE.COSIGN_ARCH_PACKAGE_SIGNATURE
            MODULE._COSIGN_PROVENANCE_CACHE.clear()
            try:
                MODULE.COSIGN_ARCH_PACKAGE = root / "missing-cosign.pkg.tar.zst"
                MODULE.COSIGN_ARCH_PACKAGE_SIGNATURE = root / "missing-cosign.pkg.tar.zst.sig"
                blocked = MODULE.verify_signed_verification_summary(
                    statement_path, bundle, subject, profile, policy_path, "release-signer"
                )
                self.assertEqual(blocked["status"], "FAIL")
                self.assertEqual(blocked["authenticity"], "NOT_VERIFIED")
                self.assertEqual(blocked["tool"]["provenance"]["status"], "FAIL")
                self.assertTrue(any("cosign verifier" in item for item in blocked["failures"]))
            finally:
                MODULE.COSIGN_ARCH_PACKAGE = original_package
                MODULE.COSIGN_ARCH_PACKAGE_SIGNATURE = original_signature
                MODULE._COSIGN_PROVENANCE_CACHE.clear()

    def test_signed_vsa_rejects_detached_statement_substitution(self) -> None:
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            subject = root / "artifact.bin"
            subject.write_bytes(b"artifact")
            profile = ROOT / "artifact-delivery/examples/pdu-sdu-presentation-r1.json"
            statement_path = root / "structural.vsa.json"
            statement = MODULE.verification_summary_statement(
                subject, profile, MODULE.LOCAL_VSA_VERIFIER_ID, {"test": "1"}, True
            )
            MODULE.write_json(statement_path, statement)
            private_key, policy_path, signing_config, env = self._cosign_key_and_policy(root)
            bundle = root / "structural.sigstore.json"
            self._sign_vsa_no_tlog(subject, statement_path, bundle, private_key, signing_config, env)
            statement["predicate"]["verifier"]["version"]["test"] = "substituted"
            MODULE.write_json(statement_path, statement)
            result = MODULE.verify_signed_verification_summary(
                statement_path, bundle, subject, profile, policy_path, "release-signer"
            )
            self.assertEqual(result["status"], "FAIL")
            self.assertTrue(any("bundle shape/statement" in item for item in result["failures"]))
            self.assertFalse(result["bundleShape"]["signedStatementMatches"])

    def test_signed_vsa_rejects_signer_verifier_policy_mismatch(self) -> None:
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            subject = root / "artifact.bin"
            subject.write_bytes(b"artifact")
            profile = ROOT / "artifact-delivery/examples/pdu-sdu-presentation-r1.json"
            statement_path = root / "structural.vsa.json"
            MODULE.write_json(statement_path, MODULE.verification_summary_statement(
                subject, profile, MODULE.LOCAL_VSA_VERIFIER_ID, {"test": "1"}, True
            ))
            private_key, policy_path, signing_config, env = self._cosign_key_and_policy(root)
            bundle = root / "structural.sigstore.json"
            self._sign_vsa_no_tlog(subject, statement_path, bundle, private_key, signing_config, env)
            policy = json.loads(policy_path.read_text())
            policy["signers"][0]["allowedVerifierIds"] = ["https://example.test/other-verifier"]
            policy_path.write_text(json.dumps(policy))
            result = MODULE.verify_signed_verification_summary(
                statement_path, bundle, subject, profile, policy_path, "release-signer"
            )
            self.assertEqual(result["status"], "FAIL")
            self.assertTrue(any("verifier.id is not authorized" in item for item in result["failures"]))

    def test_signed_vsa_rejects_signature_from_untrusted_key(self) -> None:
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            signer_root = root / "signer"
            trust_root = root / "trust"
            signer_root.mkdir()
            trust_root.mkdir()
            subject = root / "artifact.bin"
            subject.write_bytes(b"artifact")
            profile = ROOT / "artifact-delivery/examples/pdu-sdu-presentation-r1.json"
            statement_path = root / "structural.vsa.json"
            MODULE.write_json(statement_path, MODULE.verification_summary_statement(
                subject, profile, MODULE.LOCAL_VSA_VERIFIER_ID, {"test": "1"}, True
            ))
            private_key, _, signing_config, env = self._cosign_key_and_policy(signer_root)
            _, unrelated_policy, _, _ = self._cosign_key_and_policy(trust_root)
            bundle = root / "structural.sigstore.json"
            self._sign_vsa_no_tlog(subject, statement_path, bundle, private_key, signing_config, env)
            result = MODULE.verify_signed_verification_summary(
                statement_path, bundle, subject, profile, unrelated_policy, "release-signer"
            )
            self.assertEqual(result["status"], "FAIL")
            self.assertEqual(result["authenticity"], "NOT_VERIFIED")
            self.assertTrue(any("cryptographic attestation verification failed" in item for item in result["failures"]))

    def test_signed_vsa_rejects_legacy_bundle_shape(self) -> None:
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            statement = root / "vsa.json"
            statement.write_text(json.dumps({"_type": MODULE.IN_TOTO_STATEMENT_V1}))
            legacy = root / "legacy.json"
            legacy.write_text(json.dumps({"base64Signature": "AA==", "cert": "not-a-certificate"}))
            result = MODULE.verify_sigstore_vsa_bundle_shape(
                legacy, statement, [MODULE.SIGSTORE_BUNDLE_V03], "keyless"
            )
            self.assertEqual(result["status"], "FAIL")
            self.assertTrue(any("standardized Sigstore" in item for item in result["failures"]))

    def test_standard_bundle_cosign_safety_floor_is_3_0_6(self) -> None:
        self.assertFalse(MODULE._version_at_least((3, 0, 5), MODULE.COSIGN_STANDARD_BUNDLE_MIN_VERSION))
        self.assertTrue(MODULE._version_at_least((3, 0, 6), MODULE.COSIGN_STANDARD_BUNDLE_MIN_VERSION))
        self.assertTrue(MODULE._version_at_least((3, 1, 0), MODULE.COSIGN_STANDARD_BUNDLE_MIN_VERSION))

    def test_vsa_aggregation_separates_verification_from_assembly_gates(self) -> None:
        if importlib.util.find_spec("jsonschema") is None:
            self.skipTest("jsonschema is not available in this Python environment")
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            subject = root / "artifact.pptx"
            subject.write_bytes(b"artifact")
            profile = json.loads((ROOT / "artifact-delivery/examples/pdu-sdu-presentation-r1.json").read_text())
            profile.pop("$schema", None)
            for gate in list(profile["gates"]):
                profile["gates"][gate] = False
            profile["gates"]["profileSchema"] = True
            profile["gates"]["companionPdf"] = True
            profile_path = root / "profile.json"
            profile_path.write_text(json.dumps(profile))
            vsa = root / "profileSchema.vsa.json"
            vsa.write_text(json.dumps(MODULE.verification_summary_statement(
                subject, profile_path, MODULE.LOCAL_VSA_VERIFIER_ID, {"jsonschema": "4.26.0"}, True
            )))
            result = MODULE.aggregate_vsa_gates(
                profile_path, subject, {"profileSchema": vsa}, allow_local_unsigned=True
            )
            self.assertEqual(result["status"], "PASS", result)
            self.assertEqual(result["requiredGates"], ["profileSchema"])
            self.assertEqual(result["assemblyGates"], ["companionPdf"])

    def test_pptx_package_relationships_and_placeholder_gate(self) -> None:
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            good = root / "good.pptx"
            bad = root / "bad.pptx"
            make_pptx(good)
            make_pptx(bad, "TODO replace")
            result = MODULE.inspect_pptx(good, ["TODO"])
            self.assertEqual(result["status"], "PASS")
            self.assertEqual(result["package"]["slideCount"], 1)
            self.assertIn("Microsoft YaHei", result["package"]["referencedTypefaceNames"])
            self.assertIn("Microsoft YaHei", result["package"]["renderExplicitTypefaceNames"])
            self.assertEqual(result["scope"]["OOXMLSchemaValidation"], "NOT_RUN")
            bad_result = MODULE.inspect_pptx(bad, ["TODO"])
            self.assertEqual(bad_result["status"], "FAIL")
            self.assertEqual(len(bad_result["package"]["placeholderHits"]), 1)

    def test_presentation_semantics_enforce_aspect_ratio_and_slide_count(self) -> None:
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            good = root / "good.pptx"
            make_pptx(good)
            inspected = MODULE.inspect_pptx(good)
            profile = json.loads((ROOT / "artifact-delivery/examples/pdu-sdu-presentation-r1.json").read_text())
            semantic = MODULE.verify_presentation_semantics(profile, inspected)
            self.assertEqual(semantic["status"], "PASS")
            self.assertAlmostEqual(semantic["observedAspectRatio"], 16/9, places=3)
            profile["semanticPolicy"]["minimumSlideCount"] = 2
            self.assertEqual(MODULE.verify_presentation_semantics(profile, inspected)["status"], "FAIL")

            bad = root / "tall.pptx"
            tall_presentation = PRESENTATION_XML.replace("cx='12192000' cy='6858000'", "cx='9144000' cy='31089600'")
            with zipfile.ZipFile(bad, "w") as z:
                z.writestr("[Content_Types].xml", CONTENT_TYPES)
                z.writestr("_rels/.rels", ROOT_RELS)
                z.writestr("ppt/presentation.xml", tall_presentation)
                z.writestr("ppt/_rels/presentation.xml.rels", PRESENTATION_RELS)
                z.writestr("ppt/slides/slide1.xml", SLIDE_XML)
            profile["semanticPolicy"]["minimumSlideCount"] = 1
            tall = MODULE.verify_presentation_semantics(profile, MODULE.inspect_pptx(bad))
            self.assertEqual(tall["status"], "FAIL")
            self.assertTrue(any("aspect ratio" in item for item in tall["failures"]))

    def test_master_and_layout_typefaces_are_render_dependencies(self) -> None:
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / "master-font.pptx"
            with zipfile.ZipFile(path, "w") as z:
                z.writestr("[Content_Types].xml", CONTENT_TYPES)
                z.writestr("_rels/.rels", ROOT_RELS)
                z.writestr("ppt/presentation.xml", PRESENTATION_XML)
                z.writestr("ppt/_rels/presentation.xml.rels", PRESENTATION_RELS)
                z.writestr("ppt/slides/slide1.xml", SLIDE_XML)
                z.writestr("ppt/slideLayouts/slideLayout1.xml", SLIDE_XML.replace("Microsoft YaHei", "Aptos"))
                z.writestr("ppt/slideMasters/slideMaster1.xml", SLIDE_XML.replace("Microsoft YaHei", "Arial"))
            result = MODULE.inspect_pptx(path)
            self.assertEqual(result["status"], "PASS")
            self.assertEqual(
                result["package"]["renderExplicitTypefaceNames"],
                ["Aptos", "Arial", "Microsoft YaHei"],
            )

    def test_missing_internal_relationship_fails(self) -> None:
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / "broken.pptx"
            with zipfile.ZipFile(path, "w") as z:
                z.writestr("[Content_Types].xml", CONTENT_TYPES)
                z.writestr("_rels/.rels", ROOT_RELS)
                z.writestr("ppt/presentation.xml", PRESENTATION_XML)
                z.writestr("ppt/_rels/presentation.xml.rels", PRESENTATION_RELS)
            result = MODULE.inspect_pptx(path)
            self.assertEqual(result["status"], "FAIL")
            self.assertTrue(result["package"]["unresolvedRelationships"])

    def test_openxml_evidence_must_bind_exact_artifact_and_standard_validator(self) -> None:
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            pptx = root / "deck.pptx"
            make_pptx(pptx)
            evidence = root / "openxml.json"
            evidence.write_text(json.dumps({
                "status": "PASS",
                "artifact": {"sha256": MODULE.sha256_file(pptx)},
                "validator": {
                    "implementation": "DocumentFormat.OpenXml",
                    "packageVersion": "3.5.1",
                    "api": "OpenXmlValidator"
                },
                "validationErrorCount": 0,
                "validationErrors": []
            }))
            self.assertEqual(MODULE.verify_openxml_evidence(evidence, pptx)["status"], "PASS")
            value = json.loads(evidence.read_text())
            value["validator"]["implementation"] = "custom-validator"
            evidence.write_text(json.dumps(value))
            self.assertEqual(MODULE.verify_openxml_evidence(evidence, pptx)["status"], "FAIL")

    def test_powerpoint_worker_uses_native_pdf_and_png_exports(self) -> None:
        worker = (ROOT / "scripts/powerpoint_render_worker.ps1").read_text()
        self.assertIn("$presentation.SaveAs($stagePdf, 32)", worker)
        self.assertIn("$presentation.Export($stageRenderDir, 'PNG'", worker)
        self.assertIn("Presentation.SaveAs/ppSaveAsPDF", worker)
        self.assertNotIn("$presentation.ExportAsFixedFormat(", worker)

    def test_powerpoint_worker_digest_fences_input_for_full_target_run(self) -> None:
        worker = (ROOT / "scripts/powerpoint_render_worker.ps1").read_text()
        self.assertIn("[Parameter(Mandatory=$true)][ValidatePattern('^[0-9a-fA-F]{64}$')][string]$ExpectedSha256", worker)
        self.assertIn("[System.IO.FileShare]::Read", worker)
        self.assertIn("$artifactShaBefore = Get-StreamSha256 $sourceStream", worker)
        self.assertIn("if ($artifactShaBefore -ne $expected)", worker)
        self.assertIn("$artifactShaAfter = Get-StreamSha256 $sourceStream", worker)
        self.assertIn("Input PPTX changed during target render", worker)
        self.assertIn("persistent-read-handle", worker)
        self.assertIn("[System.IO.Path]::GetTempPath()", worker)
        self.assertIn("windows-local-ntfs-temp", worker)
        self.assertIn("PowerPoint PDF copy digest mismatch", worker)
        self.assertIn("PowerPoint PNG copy digest mismatch", worker)
        self.assertLess(worker.index("if ($artifactShaBefore -ne $expected)"), worker.index("New-Object -ComObject PowerPoint.Application"))
        self.assertLess(worker.index("$artifactShaAfter = Get-StreamSha256 $sourceStream"), worker.index("$sourceStream.Dispose()"))

    def test_openxml_project_pins_mature_sdk_and_uses_openxml_validator(self) -> None:
        csproj = (ROOT / "artifact-delivery/openxml-validator/ArtifactOpenXmlValidator.csproj").read_text()
        program = (ROOT / "artifact-delivery/openxml-validator/Program.cs").read_text()
        self.assertIn('PackageReference Include="DocumentFormat.OpenXml" Version="3.5.1"', csproj)
        self.assertIn("new OpenXmlValidator()", program)
        self.assertIn("PresentationDocument.Open", program)
        self.assertIn("WordprocessingDocument.Open", program)
        self.assertIn("SpreadsheetDocument.Open", program)
        self.assertNotIn("OrdivonOpenXml", program)

    def test_cross_format_toolchain_keeps_failed_writers_out_of_production_dependencies(self) -> None:
        plan = json.loads((ROOT / "artifact-delivery/toolchain-v1.plan.json").read_text())
        node = json.loads((ROOT / "artifact-delivery/node/package.json").read_text())
        requirements = (ROOT / "config/artifact-delivery-requirements.txt").read_text()
        self.assertEqual(plan["presentation"]["writer"]["tool"], "python-pptx")
        self.assertEqual(plan["spreadsheet"]["writer"]["tool"], "XlsxWriter")
        self.assertNotIn("pptxgenjs", node["dependencies"])
        self.assertIn("python-pptx==1.0.2", requirements)
        self.assertIn("XlsxWriter==3.2.9", requirements)
        self.assertNotIn("openpyxl==", requirements)

    def test_undeclared_artifact_typeface_fails_dependency_gate(self) -> None:
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            pptx = root / "aptos.pptx"
            make_pptx(pptx, font="Aptos")
            inspected = MODULE.inspect_pptx(pptx)
            self.assertIn("Aptos", inspected["package"]["renderExplicitTypefaceNames"])
            font_dir = root / "fonts"
            font_dir.mkdir()
            for name in ("msyh.ttc", "msyhbd.ttc", "msyhl.ttc"):
                (font_dir / name).write_bytes(name.encode())
            profile = json.loads((ROOT / "artifact-delivery/examples/pdu-sdu-presentation-r1.json").read_text())
            result = MODULE.verify_font_manifest(
                profile, font_dir, inspected["package"]["renderExplicitTypefaceNames"]
            )
            self.assertEqual(result["status"], "FAIL")
            self.assertEqual(result["undeclaredObservedTypefaces"], ["Aptos"])

    def test_font_manifest_binds_target_files_and_digests(self) -> None:
        with tempfile.TemporaryDirectory() as d:
            font_dir = Path(d)
            (font_dir / "font.ttc").write_bytes(b"font")
            profile = {"fonts": [{"family": "Test", "required": True, "targetFiles": ["font.ttc"], "embeddingPermission": "unknown"}]}
            result = MODULE.verify_font_manifest(profile, font_dir)
            self.assertEqual(result["status"], "PASS")
            self.assertEqual(len(result["fonts"][0]["files"][0]["sha256"]), 64)
            (font_dir / "font.ttc").unlink()
            self.assertEqual(MODULE.verify_font_manifest(profile, font_dir)["status"], "FAIL")

    def test_readback_is_digest_exact(self) -> None:
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            source = root / "source.pptx"
            readback = root / "readback.pptx"
            source.write_bytes(b"same")
            shutil.copyfile(source, readback)
            self.assertEqual(MODULE.verify_readback(source, readback)["status"], "PASS")
            readback.write_bytes(b"different")
            self.assertEqual(MODULE.verify_readback(source, readback)["status"], "FAIL")

    def test_slsa_attestation_uses_in_toto_statement_and_standard_predicate(self) -> None:
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            subject = root / "report.pptx"
            material = root / "narrative.md"
            subject.write_bytes(b"pptx")
            material.write_text("content")
            statement = MODULE.slsa_statement(
                [subject],
                [material],
                ROOT / "artifact-delivery/examples/pdu-sdu-presentation-r1.json",
                "https://builder.example.test/artifact-delivery",
                "https://builder.example.test/artifact-delivery/presentation/v1",
            )
            self.assertEqual(statement["_type"], "https://in-toto.io/Statement/v1")
            self.assertEqual(statement["predicateType"], "https://slsa.dev/provenance/v1")
            self.assertEqual(statement["subject"][0]["digest"]["sha256"], MODULE.sha256_file(subject))
            self.assertIn("buildDefinition", statement["predicate"])
            self.assertIn("runDetails", statement["predicate"])

    def test_presentation_gate_fails_closed_on_unimplemented_required_gate(self) -> None:
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            profile_path = root / "profile.json"
            profile = json.loads((ROOT / "artifact-delivery/examples/pdu-sdu-presentation-r1.json").read_text())
            profile.pop("$schema", None)
            profile["gates"]["accessibility"] = True
            profile_path.write_text(json.dumps(profile))
            pptx = root / "deck.pptx"
            make_pptx(pptx)
            font_dir = root / "fonts"
            font_dir.mkdir()
            for name in ("msyh.ttc", "msyhbd.ttc", "msyhl.ttc"):
                (font_dir / name).write_bytes(name.encode())
            result = MODULE.presentation_gate(profile_path, pptx, None, None, None, None, None, [], font_dir)
            self.assertEqual(result["status"], "FAIL")
            self.assertIn("accessibility", result["requiredGateFailures"])

    def test_toolchain_web_plan_does_not_false_green_webkit(self) -> None:
        plan = json.loads((ROOT / "artifact-delivery/toolchain-v1.plan.json").read_text())
        renderers = {item["tool"]: item for item in plan["web"]["renderers"]}
        self.assertEqual(renderers["Chromium"]["status"], "HEADLESS_LAUNCH_AND_AXE_SMOKE_PASS")
        self.assertEqual(renderers["Firefox"]["status"], "HEADLESS_LAUNCH_SMOKE_PASS")
        self.assertEqual(renderers["WebKit"]["status"], "LOCAL_ARCH_ABI_BLOCKED")
        self.assertNotIn("PASS", renderers["WebKit"]["status"])

    def test_presentation_gate_fails_closed_without_target_and_visual_evidence(self) -> None:
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            pptx = root / "deck.pptx"
            make_pptx(pptx)
            font_dir = root / "fonts"
            font_dir.mkdir()
            for name in ("msyh.ttc", "msyhbd.ttc", "msyhl.ttc"):
                (font_dir / name).write_bytes(name.encode())
            result = MODULE.presentation_gate(
                ROOT / "artifact-delivery/examples/pdu-sdu-presentation-r1.json",
                pptx,
                None,
                None,
                None,
                None,
                None,
                [],
                font_dir,
            )
            self.assertEqual(result["status"], "FAIL")
            self.assertIn("structural", result["requiredGateFailures"])
            self.assertIn("target", result["requiredGateFailures"])
            self.assertIn("visual", result["requiredGateFailures"])
            self.assertIn("companionPdf", result["requiredGateFailures"])
            self.assertIn("deliveryReadback", result["requiredGateFailures"])

    def test_verapdf_rejects_ordinary_powerpoint_pdf_as_pdfua2_when_available(self) -> None:
        pdf = ROOT / ".cache/artifact-toolchain/python-pptx/target/probe.pdf"
        if not pdf.is_file() or MODULE._verapdf_executable() is None:
            self.skipTest("local veraPDF/PDF probe is not available")
        result = MODULE.verify_pdf_conformance(pdf, "ua2")
        self.assertEqual(result["status"], "FAIL")
        self.assertFalse(result["compliant"])
        self.assertEqual(result["validator"]["implementation"], "veraPDF")

    def test_render_evidence_accepts_powerpoint_uppercase_png_extension(self) -> None:
        with tempfile.TemporaryDirectory() as d:
            render_dir = Path(d) / "renders"
            render_dir.mkdir()
            (render_dir / "幻灯片1.PNG").write_bytes(b"png")
            result = MODULE.verify_render_evidence(render_dir, 1)
            self.assertEqual(result["status"], "PASS")
            self.assertEqual(result["digests"][0]["name"], "幻灯片1.PNG")

    def test_target_evidence_binds_exact_render_digests(self) -> None:
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            pptx = root / "deck.pptx"
            pdf = root / "deck.pdf"
            make_pptx(pptx)
            pdf.write_bytes(b"pdf")
            render_dir = root / "renders"
            render_dir.mkdir()
            render = render_dir / "Slide1.PNG"
            render.write_bytes(b"png")
            render_result = MODULE.verify_render_evidence(render_dir, 1)
            evidence = root / "target.json"
            evidence.write_text(json.dumps({
                "artifact": {"sha256": MODULE.sha256_file(pptx)},
                "renderer": {"name": "Microsoft PowerPoint Desktop", "version": "16.0"},
                "result": {
                    "slideCount": 1,
                    "pdfSha256": MODULE.sha256_file(pdf),
                    "pngCount": 1,
                    "pngs": [{"name": "Slide1.PNG", "sha256": MODULE.sha256_file(render)}]
                }
            }))
            self.assertEqual(MODULE.verify_target_evidence(evidence, pptx, pdf, 1, render_result)["status"], "PASS")
            value = json.loads(evidence.read_text())
            value["result"]["pngs"][0]["sha256"] = "0" * 64
            evidence.write_text(json.dumps(value))
            self.assertEqual(MODULE.verify_target_evidence(evidence, pptx, pdf, 1, render_result)["status"], "FAIL")

    def test_visual_review_binds_exact_artifact_and_render_digests(self) -> None:
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            pptx = root / "deck.pptx"
            make_pptx(pptx)
            render_dir = root / "renders"
            render_dir.mkdir()
            (render_dir / "Slide1.png").write_bytes(b"png")
            renders = MODULE.verify_render_evidence(render_dir, 1)
            evidence = root / "visual.json"
            evidence.write_text(json.dumps({
                "artifactSha256": MODULE.sha256_file(pptx),
                "verdict": "PASS",
                "blockingDefects": [],
                "methods": ["agent-vision-review"],
                "renderDigests": {item["name"]: item["sha256"] for item in renders["digests"]},
            }))
            self.assertEqual(MODULE.verify_visual_review(evidence, pptx, renders)["status"], "PASS")
            value = json.loads(evidence.read_text())
            value["artifactSha256"] = "0" * 64
            evidence.write_text(json.dumps(value))
            self.assertEqual(MODULE.verify_visual_review(evidence, pptx, renders)["status"], "FAIL")

    def test_delivery_evidence_requires_every_target_and_required_output(self) -> None:
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            pptx = root / "deck.pptx"
            pdf = root / "deck.pdf"
            pptx.write_bytes(b"pptx")
            pdf.write_bytes(b"pdf")
            profile = json.loads((ROOT / "artifact-delivery/examples/pdu-sdu-presentation-r1.json").read_text())
            paths = []
            for destination in ("google-drive", "windows-workstation"):
                for role, source in (("primary", pptx), ("companion", pdf)):
                    readback = root / f"{destination}-{role}.bin"
                    shutil.copyfile(source, readback)
                    evidence = MODULE.verify_readback(
                        source,
                        readback,
                        destination,
                        f"ref:{destination}:{role}",
                        role,
                    )
                    evidence_path = root / f"{destination}-{role}.json"
                    evidence_path.write_text(json.dumps(evidence))
                    paths.append(evidence_path)
            self.assertEqual(MODULE.verify_delivery_evidence(paths, profile, pptx, pdf)["status"], "PASS")
            self.assertEqual(MODULE.verify_delivery_evidence(paths[:-1], profile, pptx, pdf)["status"], "FAIL")


    def _fake_pandoc(self, root: Path) -> Path:
        path = root / "pandoc"
        path.write_text("""#!/usr/bin/env python3
import json,sys
from pathlib import Path
if '--version' in sys.argv:
    print('pandoc 3.10.2')
    raise SystemExit(0)
fmt=sys.argv[sys.argv.index('--from')+1]
source=Path(sys.argv[-1])
text='different' if fmt=='docx' and source.read_bytes()==b'drift' else 'same content'
def s(value): return {'t':'Str','c':value}
ast={'pandoc-api-version':[1,23,1], 'meta':{'title':{'t':'MetaInlines','c':[s('Report')]}}, 'blocks':[{'t':'Header','c':[1,['',[],[]],[s('Heading')]]},{'t':'Para','c':[s(text)]},{'t':'OrderedList','c':[[1,{'t':'Decimal'},{'t':'Period'}],[[{'t':'Plain','c':[s('item')]}]]]}]}
print(json.dumps(ast))
""")
        path.chmod(0o755)
        return path

    def test_document_semantic_correspondence_is_digest_bound_and_detects_drift(self) -> None:
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            pandoc = self._fake_pandoc(root)
            source = root / "source.md"
            document = root / "artifact.docx"
            source.write_text("# Heading\n\nsame content\n\n1. item\n")
            document.write_bytes(b"same")
            result = MODULE.verify_document_semantic_correspondence(source, document, pandoc)
            self.assertEqual(result["status"], "PASS", result)
            self.assertEqual(result["sourceProjection"], result["documentProjection"])
            self.assertEqual(result["metadata"]["title"]["matched"], True)
            self.assertIn("not independent IV&V", result["boundary"])
            document.write_bytes(b"drift")
            drift = MODULE.verify_document_semantic_correspondence(source, document, pandoc)
            self.assertEqual(drift["status"], "FAIL", drift)
            self.assertTrue(any("semantic projections differ" in item for item in drift["failures"]))

    def test_document_dependency_verifier_binds_pandoc_archive_and_request(self) -> None:
        if importlib.util.find_spec("jsonschema") is None:
            self.skipTest("jsonschema is unavailable")
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            pandoc = self._fake_pandoc(root)
            archive = root / "pandoc.tar.gz"
            archive.write_bytes(b"official-release-archive")
            validator = root / "validate-openxml"
            validator.write_text("#!/bin/sh\nexit 0\n")
            validator.chmod(0o755)
            profile = root / "document-r1.json"
            shutil.copyfile(ROOT / "artifact-delivery/examples/document-r1.json", profile)
            source = root / "source.md"
            source.write_text("# Report\n")
            artifact = root / "artifact.docx"
            artifact.write_bytes(b"docx")
            lock = root / "toolchain.json"
            lock.write_text(json.dumps({"pandoc": {"version": "3.10.2", "binarySha256": MODULE.sha256_file(pandoc), "linuxAmd64ArchiveSha256": MODULE.sha256_file(archive)}}))
            request = root / "request.json"
            request.write_text(json.dumps({
                "schemaVersion": 1,
                "kind": "artifact-delivery-request",
                "requestId": "artifact-request:document-dependency-test",
                "profile": {"id": "document-r1", "path": profile.name, "sha256": MODULE.sha256_file(profile)},
                "source": {"kind": "markdown", "path": source.name, "sha256": MODULE.sha256_file(source)},
                "materials": [],
                "outputDirectory": "out",
                "builder": {"id": "https://ordivon.local/builders/artifact-delivery/pandoc-v1", "buildType": "https://ordivon.local/build-types/artifact-delivery/markdown-docx-v1"}
            }))
            result = MODULE.verify_document_dependencies(request, artifact, pandoc, archive, lock, validator)
            self.assertEqual(result["status"], "PASS", result)
            self.assertTrue(result["pandoc"]["digestMatched"])
            self.assertTrue(result["pandocReleaseArchive"]["digestMatched"])
            archive.write_bytes(b"changed")
            drift = MODULE.verify_document_dependencies(request, artifact, pandoc, archive, lock, validator)
            self.assertEqual(drift["status"], "FAIL", drift)
            self.assertTrue(any("release archive digest" in item for item in drift["failures"]))


if __name__ == "__main__":
    unittest.main()
